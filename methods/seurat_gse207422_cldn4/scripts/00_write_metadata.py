#!/usr/bin/env python3
"""Write GSE207422 scRNA sample metadata as TSV (ICI labels included)."""
from __future__ import annotations

from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parents[1]
REPO = HERE.parents[1]
DEFAULT_XLSX = REPO / "data/GSE207422/GSE207422_NSCLC_scRNAseq_metadata.xlsx"


def main() -> None:
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--out", type=Path, default=HERE / "sample_metadata.tsv")
    args = ap.parse_args()
    xlsx = args.xlsx
    out = args.out
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header = [str(x) if x is not None else "" for x in rows[0]]
    keep = []
    for raw in rows[1:]:
        if not raw or raw[0] is None:
            continue
        sample = str(raw[0])
        if sample.startswith("RECIST") or sample.startswith("MPR") or sample.startswith("NMPR") or sample.startswith("pCR"):
            continue
        rec = {header[i]: raw[i] if i < len(raw) else None for i in range(len(header))}
        resource = str(rec.get("Resource") or "")
        timing = "pre" if "Pre" in resource else "post"
        path_resp = str(rec.get("Pathologic Response") or "")
        ici_group = "MPR" if path_resp in {"MPR", "pCR"} else path_resp
        keep.append(
            [
                sample,
                rec.get("Patient"),
                resource,
                timing,
                rec.get("Sex"),
                rec.get("Age"),
                rec.get("Clinical Stage"),
                rec.get("Pathology"),
                rec.get("PD1 Antibody"),
                rec.get("Chemotherapy"),
                path_resp,
                ici_group,
                rec.get("Residual Tumor"),
                rec.get("RECIST"),
            ]
        )
    cols = [
        "Sample",
        "Patient",
        "Resource",
        "timing",
        "Sex",
        "Age",
        "Clinical_Stage",
        "Pathology",
        "PD1_Antibody",
        "Chemotherapy",
        "Pathologic_Response",
        "ici_group",
        "Residual_Tumor",
        "RECIST",
    ]
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w") as fh:
        fh.write("\t".join(cols) + "\n")
        for r in keep:
            fh.write("\t".join("" if x is None else str(x) for x in r) + "\n")
    print(f"wrote {out} n={len(keep)}", flush=True)


if __name__ == "__main__":
    main()
