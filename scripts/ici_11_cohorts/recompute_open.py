#!/usr/bin/env python3
"""Recompute CLDN4 / TACSTD2 on the open pieces of the Lee 2024 11 ICI cohorts.

Source of the cohort list: Lee et al., Sci Adv 2024 (PMID 38295179), Data S1.
Response labels for the nine bulk cohorts: their Data S9 (CR/PR = responder,
SD/PD = non-responder), committed as data/ici_11_cohorts/lee2024_data_s9_labels.tsv.

Bessede et al., Clin Cancer Res 2024 (PMID 38048058) is the citable TROP2 (TACSTD2)
ICI paper. It is not a CLDN4 meta and its POPLAR/OAK/BIP matrices are not open.
This script does not impute those cohorts.
"""

from __future__ import annotations

import gzip
import json
import os
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from scipy.stats import mannwhitneyu, spearmanr

ROOT = Path(__file__).resolve().parents[2]
LABELS = ROOT / "data/ici_11_cohorts/lee2024_data_s9_labels.tsv"
CACHE = Path(os.environ.get("ICI11_CACHE", "/tmp/ici11"))
OUT = ROOT / "results/ici_11_cohorts"
CACHE.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)

ENTREZ = {"CLDN4": 1364, "TACSTD2": 4070, "CD8A": 925}
ENSG = {
    "CLDN4": "ENSG00000189143",
    "TACSTD2": "ENSG00000184292",
    "CD8A": "ENSG00000153563",
}


def fetch_json(url: str, payload: dict | None = None):
    data = None if payload is None else json.dumps(payload).encode()
    headers = {"Accept": "application/json", "User-Agent": "Mozilla/5.0"}
    if payload is not None:
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(url, data=data, headers=headers)
    with urllib.request.urlopen(req, timeout=180) as resp:
        return json.load(resp)


def download(url: str, dest: Path) -> Path:
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    urllib.request.urlretrieve(url, dest)
    return dest


def load_labels() -> pd.DataFrame:
    df = pd.read_csv(LABELS, sep="\t")
    df["y"] = df["response"].map({"responder": 1, "non-responder": 0})
    if df["y"].isna().any():
        raise SystemExit(f"unmapped labels: {df.loc[df.y.isna(), 'response'].unique()}")
    return df


def series_table(path: Path) -> pd.DataFrame:
    gsm, title = [], []
    with gzip.open(path, "rt", errors="replace") as handle:
        for line in handle:
            if line.startswith("!Sample_geo_accession"):
                gsm = [x.strip().strip('"') for x in line.rstrip("\n").split("\t")[1:]]
            elif line.startswith("!Sample_title"):
                title = [x.strip().strip('"') for x in line.rstrip("\n").split("\t")[1:]]
    if not gsm or len(gsm) != len(title):
        raise SystemExit(f"series matrix parse failed: {path}")
    return pd.DataFrame({"sample_id": gsm, "title": title})


def cbio_genes(profile: str, sample_list: str) -> pd.DataFrame:
    mol = fetch_json(
        f"https://www.cbioportal.org/api/molecular-profiles/{profile}/molecular-data/fetch",
        {"entrezGeneIds": list(ENTREZ.values()), "sampleListId": sample_list},
    )
    gene_of = {v: k for k, v in ENTREZ.items()}
    rows = []
    for rec in mol:
        rows.append(
            {
                "sample_id": rec["sampleId"],
                "patient_id": rec.get("patientId"),
                "gene": gene_of[rec["entrezGeneId"]],
                "value": rec["value"],
            }
        )
    wide = pd.DataFrame(rows).pivot_table(
        index=["sample_id", "patient_id"], columns="gene", values="value", aggfunc="first"
    )
    return wide.reset_index()


def test_gene(values: pd.Series, y: pd.Series) -> dict:
    mask = values.notna() & y.notna()
    v = values[mask].astype(float)
    yy = y[mask].astype(int)
    r = v[yy == 1]
    nr = v[yy == 0]
    out = {
        "n": int(mask.sum()),
        "n_responder": int((yy == 1).sum()),
        "n_nonresponder": int((yy == 0).sum()),
    }
    if len(r) < 2 or len(nr) < 2 or v.nunique() < 2:
        out.update(
            {
                "auc_nr_gt_r": None,
                "mwu_p": None,
                "spearman_rho_vs_responder": None,
                "spearman_p": None,
                "median_responder": None,
                "median_nonresponder": None,
                "delta_median_nr_minus_r": None,
            }
        )
        return out
    stat = mannwhitneyu(nr, r, alternative="two-sided")
    auc = stat.statistic / (len(nr) * len(r))
    rho = spearmanr(v, yy)
    out.update(
        {
            "auc_nr_gt_r": float(auc),
            "mwu_p": float(stat.pvalue),
            "spearman_rho_vs_responder": float(rho.statistic),
            "spearman_p": float(rho.pvalue),
            "median_responder": float(r.median()),
            "median_nonresponder": float(nr.median()),
            "delta_median_nr_minus_r": float(nr.median() - r.median()),
        }
    )
    return out


def long_tests(cohort: str, frame: pd.DataFrame, genes: list[str], scale: str, note: str) -> list[dict]:
    rows = []
    for gene in genes:
        if gene not in frame.columns:
            rows.append(
                {
                    "cohort": cohort,
                    "gene": gene,
                    "gene_present": False,
                    "scale": scale,
                    "note": note,
                    "n": 0,
                }
            )
            continue
        stats = test_gene(frame[gene], frame["y"])
        stats.update(
            {
                "cohort": cohort,
                "gene": gene,
                "gene_present": True,
                "scale": scale,
                "note": note,
            }
        )
        rows.append(stats)
    return rows


def samples_from_wide(cohort: str, frame: pd.DataFrame, genes: list[str]) -> pd.DataFrame:
    keep = ["sample_id", "y", "response"] + [g for g in genes if g in frame.columns]
    out = frame[keep].copy()
    out.insert(0, "cohort", cohort)
    return out


def hugo(labels: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    xlsx = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE78nnn/GSE78220/suppl/GSE78220_PatientFPKM.xlsx",
        CACHE / "GSE78220_PatientFPKM.xlsx",
    )
    series = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE78nnn/GSE78220/matrix/GSE78220_series_matrix.txt.gz",
        CACHE / "GSE78220_series_matrix.txt.gz",
    )
    meta = series_table(series)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["FPKM"]
    header = list(next(ws.iter_rows(max_row=1, values_only=True)))
    cols = {name: i for i, name in enumerate(header)}
    wanted = {}
    for _, rec in meta.iterrows():
        title = rec["title"]
        col = None
        for candidate in (f"{title}.baseline", f"{title}.OnTx", title):
            if candidate in cols:
                col = candidate
                break
        if col:
            wanted[rec["sample_id"]] = col
    genes = {g: [] for g in ("CLDN4", "TACSTD2", "CD8A")}
    gene_rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] in genes:
            gene_rows[row[0]] = row
    lab = labels[labels.cohort == "Hugo"][["sample_id", "y", "response"]]
    records = []
    n_ontx = 0
    for _, rec in lab.iterrows():
        col = wanted.get(rec["sample_id"])
        if col is None:
            continue
        if col.endswith(".OnTx"):
            n_ontx += 1
        item = {"sample_id": rec["sample_id"], "y": rec["y"], "response": rec["response"]}
        idx = cols[col]
        for gene, grow in gene_rows.items():
            item[gene] = grow[idx]
        records.append(item)
    note = (
        f"GEO GSE78220 FPKM, Data S9 GSM labels. "
        f"{n_ontx} joined column(s) are OnTx rather than baseline."
    )
    return pd.DataFrame(records), note


def symbol_counts(path: Path, genes: list[str]) -> dict[str, dict[str, float]]:
    out = {g: {} for g in genes}
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        samples = header[1:]
        for line in handle:
            gene, *vals = line.rstrip("\n").split("\t")
            if gene in out:
                out[gene] = {s: float(v) if v not in ("", "NA") else float("nan") for s, v in zip(samples, vals)}
    return out


def jung(labels: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    expr = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE135nnn/GSE135222/suppl/GSE135222_GEO_RNA-seq_omicslab_exp.tsv.gz",
        CACHE / "GSE135222_GEO_RNA-seq_omicslab_exp.tsv.gz",
    )
    series = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE135nnn/GSE135222/matrix/GSE135222_series_matrix.txt.gz",
        CACHE / "GSE135222_series_matrix.txt.gz",
    )
    meta = series_table(series)
    meta["col"] = meta["title"].str.replace(" ", "", regex=False)
    expr_map = {}
    with gzip.open(expr, "rt") as handle:
        header = handle.readline().rstrip("\n").split("\t")
        samples = header[1:]
        want = set(ENSG.values())
        found = {}
        for line in handle:
            gid = line.split("\t", 1)[0].split(".")[0]
            if gid in want:
                vals = line.rstrip("\n").split("\t")[1:]
                found[gid] = {s: float(v) for s, v in zip(samples, vals)}
    gene_of = {v: k for k, v in ENSG.items()}
    for ens, table in found.items():
        expr_map[gene_of[ens]] = table
    lab = labels[labels.cohort == "Jung"]
    records = []
    for _, rec in lab.iterrows():
        col = meta.loc[meta.sample_id == rec["sample_id"], "col"]
        if col.empty:
            continue
        col = col.iloc[0]
        item = {"sample_id": rec["sample_id"], "y": rec["y"], "response": rec["response"]}
        for gene, table in expr_map.items():
            item[gene] = table.get(col)
        records.append(item)
    return pd.DataFrame(records), "GEO GSE135222 supplied expression (Ensembl), Data S9 GSM labels."


def cho(labels: pd.DataFrame) -> tuple[pd.DataFrame, str]:
    expr = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE126nnn/GSE126044/suppl/GSE126044_counts.txt.gz",
        CACHE / "GSE126044_counts.txt.gz",
    )
    series = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE126nnn/GSE126044/matrix/GSE126044_series_matrix.txt.gz",
        CACHE / "GSE126044_series_matrix.txt.gz",
    )
    meta = series_table(series)
    meta["col"] = meta["title"].str.replace("RNA-seq_", "", regex=False)
    tables = symbol_counts(expr, ["CLDN4", "TACSTD2", "CD8A"])
    lab = labels[labels.cohort == "Cho"]
    records = []
    for _, rec in lab.iterrows():
        col = meta.loc[meta.sample_id == rec["sample_id"], "col"]
        if col.empty:
            continue
        col = col.iloc[0]
        item = {"sample_id": rec["sample_id"], "y": rec["y"], "response": rec["response"]}
        for gene, table in tables.items():
            item[gene] = table.get(col)
        records.append(item)
    return pd.DataFrame(records), (
        "GEO GSE126044 raw counts joined to Data S9 GSM labels. "
        "GEO 'patient response' disagrees for GSM3589680 (Dis_17): GEO responder, Data S9 non-responder. "
        "That sample has high CLDN4 and high CD8A. Using the GEO string for that one sample "
        "moves the CLDN4 AUC from the Data S9 value to about 0.73 (two-sided p about 0.18)."
    )


def liu_vanallen(labels: pd.DataFrame, cohort: str, profile: str, sample_list: str, note: str) -> pd.DataFrame:
    wide = cbio_genes(profile, sample_list)
    lab = labels[labels.cohort == cohort][["sample_id", "y", "response"]].rename(
        columns={"sample_id": "patient_id"}
    )
    merged = wide.merge(lab, on="patient_id", how="inner")
    if merged["sample_id"].duplicated().any():
        # Prefer the sample whose id echoes the patient number.
        merged["prefer"] = merged.apply(
            lambda r: str(r["sample_id"]).replace("Sample", "Patient") == str(r["patient_id"])
            or str(r["sample_id"]) == str(r["patient_id"]),
            axis=1,
        )
        merged = merged.sort_values("prefer", ascending=False).drop_duplicates("patient_id")
    merged = merged.rename(columns={"sample_id": "expr_sample_id", "patient_id": "sample_id"})
    merged.attrs["note"] = note + f" Joined {len(merged)} Data S9 patients."
    return merged


def gide() -> pd.DataFrame:
    wide = cbio_genes(
        "mel_iatlas_gide_2019_rna_seq_mrna", "mel_iatlas_gide_2019_all"
    )
    clin = fetch_json(
        "https://www.cbioportal.org/api/studies/mel_iatlas_gide_2019/clinical-data"
        "?clinicalDataType=SAMPLE&projection=SUMMARY&pageSize=20000&pageNumber=0"
    )
    by = defaultdict(dict)
    for rec in clin:
        if rec["clinicalAttributeId"] in ("RESPONSE", "SAMPLE_TREATMENT"):
            by[rec["sampleId"]][rec["clinicalAttributeId"]] = rec["value"]
    meta = pd.DataFrame(
        [{"sample_id": k, **v} for k, v in by.items()]
    )
    frame = wide.merge(meta, on="sample_id", how="inner")
    recist = {
        "Complete Response": 1,
        "Partial Response": 1,
        "Stable Disease": 0,
        "Progressive Disease": 0,
    }
    frame["y"] = frame["RESPONSE"].map(recist)
    frame["response"] = frame["y"].map({1: "responder", 0: "non-responder"})
    frame = frame.dropna(subset=["y"])
    return frame


def imvigor() -> pd.DataFrame:
    wide = cbio_genes(
        "blca_iatlas_imvigor210_2017_rna_seq_mrna",
        "blca_iatlas_imvigor210_2017_all",
    )
    clin = fetch_json(
        "https://www.cbioportal.org/api/studies/blca_iatlas_imvigor210_2017/clinical-data"
        "?clinicalDataType=SAMPLE&projection=SUMMARY&pageSize=20000&pageNumber=0"
    )
    resp = {
        rec["sampleId"]: rec["value"]
        for rec in clin
        if rec["clinicalAttributeId"] == "RESPONSE"
    }
    frame = wide.copy()
    frame["RESPONSE"] = frame["sample_id"].map(resp)
    recist = {
        "Complete Response": 1,
        "Partial Response": 1,
        "Stable Disease": 0,
        "Progressive Disease": 0,
    }
    frame["y"] = frame["RESPONSE"].map(recist)
    frame["response"] = frame["y"].map({1: "responder", 0: "non-responder"})
    return frame.dropna(subset=["y"])


def prat_genes_present() -> list[str]:
    path = download(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE93nnn/GSE93157/suppl/GSE93157_raw_data_values.txt.gz",
        CACHE / "GSE93157_raw_data_values.txt.gz",
    )
    found = []
    with gzip.open(path, "rt", errors="replace") as handle:
        for line in handle:
            if line.startswith("#") or line.startswith("ID_REF"):
                continue
            gene = line.split("\t", 1)[0]
            if gene in ("CLDN4", "TACSTD2", "CD8A"):
                found.append(gene)
    return found


def scan_gene_row(path: Path, genes: set[str]) -> dict[str, dict]:
    """Return per-gene n_cells, n_nonzero, max for a genes-by-cells or headered matrix."""
    if not path.exists():
        return {}
    opener = gzip.open if str(path).endswith(".gz") else open
    found = {}
    with opener(path, "rt", errors="replace") as handle:
        header = handle.readline()
        if not header:
            return {}
        # genes as rows
        for line in handle:
            gene = line.split("\t", 1)[0].split(",")[0].strip().strip('"')
            gene = gene.split(".")[0]
            if gene not in genes and gene.split("|")[0] not in genes:
                continue
            key = gene if gene in genes else gene.split("|")[0]
            parts = line.rstrip("\n").replace(",", "\t").split("\t")[1:]
            vals = []
            for part in parts:
                try:
                    vals.append(float(part))
                except ValueError:
                    continue
            if not vals:
                continue
            nz = sum(v > 0 for v in vals)
            found[key] = {
                "n_values": len(vals),
                "n_nonzero": nz,
                "max": max(vals),
                "median": float(pd.Series(vals).median()),
            }
            if len(found) == len(genes):
                break
    return found


def main() -> None:
    labels = load_labels()
    tests = []
    per_sample = []

    hugo_df, hugo_note = hugo(labels)
    tests += long_tests("Hugo", hugo_df, ["CLDN4", "TACSTD2", "CD8A"], "FPKM", hugo_note)
    per_sample.append(samples_from_wide("Hugo", hugo_df, ["CLDN4", "TACSTD2", "CD8A"]))

    jung_df, jung_note = jung(labels)
    tests += long_tests(
        "Jung", jung_df, ["CLDN4", "TACSTD2", "CD8A"], "supplied expression", jung_note
    )
    per_sample.append(samples_from_wide("Jung", jung_df, ["CLDN4", "TACSTD2", "CD8A"]))

    cho_df, cho_note = cho(labels)
    tests += long_tests("Cho", cho_df, ["CLDN4", "TACSTD2", "CD8A"], "raw counts", cho_note)
    per_sample.append(samples_from_wide("Cho", cho_df, ["CLDN4", "TACSTD2", "CD8A"]))

    liu = liu_vanallen(
        labels,
        "Liu",
        "mel_dfci_2019_mrna_seq_tpm",
        "mel_dfci_2019_all",
        "cBioPortal mel_dfci_2019 TPM joined to Data S9 on patient id.",
    )
    tests += long_tests("Liu", liu, ["CLDN4", "TACSTD2", "CD8A"], "TPM", liu.attrs["note"])
    per_sample.append(samples_from_wide("Liu", liu, ["CLDN4", "TACSTD2", "CD8A"]))

    van = liu_vanallen(
        labels,
        "VanAllen",
        "skcm_dfci_2015_rna_seq_mrna",
        "skcm_dfci_2015_all",
        "cBioPortal skcm_dfci_2015 RNA-seq joined to Data S9 on patient id.",
    )
    tests += long_tests(
        "VanAllen", van, ["CLDN4", "TACSTD2", "CD8A"], "cBio continuous RNA", van.attrs["note"]
    )
    per_sample.append(samples_from_wide("VanAllen", van, ["CLDN4", "TACSTD2", "CD8A"]))

    gide_all = gide()
    gide_pre = gide_all[gide_all["SAMPLE_TREATMENT"] == "Pre"].copy()
    note_all = (
        "Open iAtlas/cBio mel_iatlas_gide_2019 log2 UQ counts. "
        "Data S9 uses ENA ERR run ids (PRJEB23709 FASTQ) with no open processed matrix at those ids. "
        "RECIST CR/PR vs SD/PD on all 91 pre+on samples matches Data S9 counts "
        f"({int((gide_all.y==1).sum())} responder / {int((gide_all.y==0).sum())} non-responder)."
    )
    note_pre = "Same Gide matrix, pretreatment biopsies only."
    tests += long_tests("Gide_all", gide_all, ["CLDN4", "TACSTD2", "CD8A"], "log2 UQ counts", note_all)
    tests += long_tests("Gide_pre", gide_pre, ["CLDN4", "TACSTD2", "CD8A"], "log2 UQ counts", note_pre)
    per_sample.append(samples_from_wide("Gide_all", gide_all, ["CLDN4", "TACSTD2", "CD8A"]))

    imv = imvigor()
    imv_note = (
        "Open iAtlas/cBio blca_iatlas_imvigor210_2017. "
        f"RESPONSE-labeled n={len(imv)} "
        f"({int((imv.y==1).sum())} CR/PR, {int((imv.y==0).sum())} SD/PD), "
        "matching Data S9 Mariathasan 68/230. Sample ids are iAtlas numbers, not the SAM ids in Data S9."
    )
    tests += long_tests(
        "Mariathasan", imv, ["CLDN4", "TACSTD2", "CD8A"], "cBio TPM-profile values", imv_note
    )
    per_sample.append(samples_from_wide("Mariathasan", imv, ["CLDN4", "TACSTD2", "CD8A"]))

    prat_found = prat_genes_present()
    tests.append(
        {
            "cohort": "PratMelanoma",
            "gene": "CLDN4",
            "gene_present": "CLDN4" in prat_found,
            "scale": "NanoString GPL19965",
            "note": (
                "GSE93157 is an immune-gene panel. "
                f"Genes present among CLDN4/TACSTD2/CD8A: {prat_found or 'none'}. "
                "Data S9 melanoma subset n=25. CLDN4 and TACSTD2 are absent, so no response test."
            ),
            "n": 25,
            "n_responder": int(((labels.cohort == "PratMelanoma") & (labels.y == 1)).sum()),
            "n_nonresponder": int(((labels.cohort == "PratMelanoma") & (labels.y == 0)).sum()),
        }
    )
    tests.append(
        {
            "cohort": "Kim",
            "gene": "CLDN4",
            "gene_present": None,
            "scale": "FASTQ only",
            "note": (
                "PRJEB25780 has 78 RNA-seq FASTQ runs and 110 exome runs, and no submitted analysis counts. "
                "Data S9 n=45 (12 responder / 33 non-responder). Not recomputed."
            ),
            "n": 45,
            "n_responder": 12,
            "n_nonresponder": 33,
        }
    )

    sade = CACHE / "GSE120575_Sade_Feldman_melanoma_single_cells_TPM_GEO.txt.gz"
    if not sade.exists():
        alt = CACHE / "GSE120575_TPM.txt.gz"
        if alt.exists():
            sade = alt
    sade_stats = scan_gene_row(sade, {"CLDN4", "TACSTD2", "CD8A"}) if sade.exists() else {}
    tests.append(
        {
            "cohort": "Sade-Feldman",
            "gene": "CLDN4",
            "gene_present": "CLDN4" in sade_stats if sade_stats else None,
            "scale": "scRNA TPM, CD45+ sorted",
            "note": (
                "GSE120575 immune-sorted single cells. Not a tumor CLDN4 response test. "
                + (
                    "Scanned gene rows: " + json.dumps(sade_stats)
                    if sade_stats
                    else "TPM matrix was not in the cache at run time; accession is open."
                )
            ),
            "n": None,
        }
    )
    tests.append(
        {
            "cohort": "Jerby-Arnon",
            "gene": "CLDN4",
            "gene_present": None,
            "scale": "scRNA TPM",
            "note": (
                "GSE115978 is open (counts, TPM, cell annotations). "
                "The cohort contrast in Lee et al. is resistant versus untreated, not pretreatment RECIST. "
                "No patient-level R/NR recompute was run."
            ),
            "n": None,
        }
    )

    test_df = pd.DataFrame(tests)
    test_df.to_csv(OUT / "open_gene_tests.tsv", sep="\t", index=False)
    pd.concat(per_sample, ignore_index=True).to_csv(OUT / "per_sample_open.tsv", sep="\t", index=False)

    # Cohort inventory, one row per Lee Data S1 cohort plus the Bessede cohorts that are not in that list.
    inventory = [
        {
            "cohort": "VanAllen",
            "cancer": "melanoma",
            "ici": "ipilimumab",
            "accession": "phs000452; open reprocess skcm_dfci_2015",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Liu",
            "cancer": "melanoma",
            "ici": "nivolumab or pembrolizumab",
            "accession": "open reprocess mel_dfci_2019 (Lee zenodo 4661265 blocked from this environment)",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Gide",
            "cancer": "melanoma",
            "ici": "anti-PD-1 +/- anti-CTLA-4",
            "accession": "PRJEB23709 FASTQ; open reprocess mel_iatlas_gide_2019",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Hugo",
            "cancer": "melanoma",
            "ici": "pembrolizumab",
            "accession": "GSE78220",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Jung",
            "cancer": "NSCLC",
            "ici": "anti-PD-1/PD-L1",
            "accession": "GSE135222",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Kim",
            "cancer": "gastric",
            "ici": "pembrolizumab",
            "accession": "PRJEB25780",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": False,
            "response_open": True,
            "cldn4_recomputed": False,
        },
        {
            "cohort": "Mariathasan",
            "cancer": "urothelial",
            "ici": "atezolizumab",
            "accession": "IMvigor210CoreBiologies; open reprocess blca_iatlas_imvigor210_2017",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Prat",
            "cancer": "melanoma (panel also has lung and HNSCC)",
            "ici": "nivolumab or pembrolizumab",
            "accession": "GSE93157",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": False,
        },
        {
            "cohort": "Cho",
            "cancer": "NSCLC",
            "ici": "nivolumab",
            "accession": "GSE126044",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": True,
        },
        {
            "cohort": "Sade-Feldman",
            "cancer": "melanoma",
            "ici": "anti-PD-1 or anti-CTLA-4",
            "accession": "GSE120575",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": True,
            "cldn4_recomputed": False,
        },
        {
            "cohort": "Jerby-Arnon",
            "cancer": "melanoma",
            "ici": "anti-PD-1 or anti-CTLA-4",
            "accession": "GSE115978",
            "in_lee2024_data_s1": True,
            "processed_matrix_open": True,
            "response_open": "resistant vs untreated, not RECIST R/NR",
            "cldn4_recomputed": False,
        },
        {
            "cohort": "POPLAR+OAK (Bessede 2024)",
            "cancer": "NSCLC",
            "ici": "atezolizumab vs docetaxel",
            "accession": "EGA EGAS00001005013; DAC EGAC00001002120",
            "in_lee2024_data_s1": False,
            "processed_matrix_open": False,
            "response_open": False,
            "cldn4_recomputed": False,
        },
        {
            "cohort": "BIP (Bessede 2024)",
            "cancer": "NSCLC",
            "ici": "ICI",
            "accession": "NCT02534649; not deposited",
            "in_lee2024_data_s1": False,
            "processed_matrix_open": False,
            "response_open": False,
            "cldn4_recomputed": False,
        },
    ]
    pd.DataFrame(inventory).to_csv(OUT / "cohort_table.tsv", sep="\t", index=False)
    print(test_df[["cohort", "gene", "n", "n_responder", "n_nonresponder", "auc_nr_gt_r", "mwu_p"]].to_string(index=False))


if __name__ == "__main__":
    main()
