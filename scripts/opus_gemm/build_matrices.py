#!/usr/bin/env python3
"""Turn the downloaded processed files into per-dataset expression matrices on a
common log2 scale, plus an auditable design table per dataset.

Outputs (under --outdir, default results/opus_gemm/data/matrices):
    <key>.expr.tsv.gz   genes x samples, log2 scale
    <key>.design.tsv    one row per sample: column, gsm, group, model, provenance
And a QC summary at --qc.

Normalisation by input type:
    counts    -> log2(CPM + 1) using each column's own library size
    tpm/fpkm/cpm -> log2(x + 1)
    vst/log2norm -> used as provided (already variance-stabilised / log2)
"""
from __future__ import annotations

import argparse
import csv
import gzip
import io
import os
import re
import sys
import tarfile

import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import geo_meta  # noqa: E402
from datasets import DATASETS, Dataset  # noqa: E402

RAW = "results/opus_gemm/data/raw"


# --------------------------------------------------------------------------- IO
def _open_text(path: str):
    return gzip.open(path, "rt") if path.endswith(".gz") else open(path)


def read_table(ds: Dataset, path: str) -> pd.DataFrame:
    df = pd.read_csv(
        path,
        sep=ds.sep,
        comment=ds.comment,
        index_col=ds.gene_col if isinstance(ds.gene_col, int) else None,
        low_memory=False,
    )
    if not isinstance(ds.gene_col, int):
        df = df.set_index(ds.gene_col)
    if ds.drop_cols:
        df = df.drop(columns=[c for c in ds.drop_cols if c in df.columns])
    return df


def read_tar_table(ds: Dataset, path: str) -> pd.DataFrame:
    series: dict[str, pd.Series] = {}
    with tarfile.open(path) as tf:
        for m in tf.getmembers():
            if not re.search(ds.member_pattern, m.name):
                continue
            blob = tf.extractfile(m).read()
            fh = io.TextIOWrapper(io.BytesIO(blob))
            if m.name.endswith(".gz"):
                fh = io.TextIOWrapper(gzip.open(io.BytesIO(blob), "rb"))
            sub = pd.read_csv(fh, sep=ds.sep, comment=ds.comment, low_memory=False)
            gene = sub.iloc[:, ds.gene_col] if isinstance(ds.gene_col, int) else sub[ds.gene_col]
            if isinstance(ds.value_col, int):
                vals = sub.iloc[:, ds.value_col]
            else:
                vals = sub[ds.value_col]
            name = os.path.basename(m.name)
            series[name] = pd.Series(vals.to_numpy(), index=gene.to_numpy(), name=name)
    if not series:
        raise RuntimeError(f"{ds.key}: no tar members matched {ds.member_pattern}")
    return pd.DataFrame(series)


def read_tar_xlsx(ds: Dataset, path: str) -> pd.DataFrame:
    import openpyxl

    series: dict[str, pd.Series] = {}
    with tarfile.open(path) as tf:
        for m in tf.getmembers():
            if not re.search(ds.member_pattern, m.name):
                continue
            blob = tf.extractfile(m).read()
            wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            ws.reset_dimensions()
            rows = ws.iter_rows(values_only=True)
            header = [str(c) if c is not None else "" for c in next(rows)]
            gi, vi = header.index(str(ds.gene_col)), header.index(str(ds.value_col))
            genes, vals = [], []
            for r in rows:
                if r[gi] is None:
                    continue
                genes.append(str(r[gi]))
                vals.append(pd.to_numeric(r[vi], errors="coerce"))
            wb.close()
            name = os.path.basename(m.name)
            series[name] = pd.Series(vals, index=genes, name=name)
    if not series:
        raise RuntimeError(f"{ds.key}: no xlsx members matched")
    return pd.DataFrame(series)


def read_xlsx(ds: Dataset, path: str) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(c) if c is not None else "" for c in rows[ds.header_row]]
    body = rows[ds.header_row + 1 :]

    def clean(v):
        if v is None:
            return ""
        s = str(v)
        m = re.match(r'^="?(.*?)"?$', s)
        return m.group(1) if m else s

    gi = header.index(str(ds.gene_col)) if not isinstance(ds.gene_col, int) else ds.gene_col
    # data columns: those matching a group rule
    keep = [i for i, h in enumerate(header) if any(re.search(pat, h) for pat, _ in ds.groups)]
    if not keep:
        raise RuntimeError(f"{ds.key}: no xlsx columns matched group rules; header={header[:40]}")
    # GSE241978 repeats sample names (raw block then normalised block); keep the last block
    seen: dict[str, int] = {}
    for i in keep:
        seen[header[i]] = i
    keep = sorted(seen.values())
    genes = [clean(r[gi]) for r in body]
    data = {header[i]: pd.to_numeric([r[i] for r in body], errors="coerce") for i in keep}
    df = pd.DataFrame(data, index=genes)
    return df[~df.index.duplicated(keep="first")]


READERS = {"table": read_table, "tar_table": read_tar_table, "tar_xlsx": read_tar_xlsx, "xlsx": read_xlsx}


# ------------------------------------------------------------------- transforms
def clean_columns(ds: Dataset, cols: list[str]) -> list[str]:
    out = []
    for c in cols:
        c = str(c)
        for pat, repl in ds.col_clean:
            c = re.sub(pat, repl, c)
        out.append(c.strip())
    return out


def collapse_genes(df: pd.DataFrame, value_type: str) -> pd.DataFrame:
    """One row per gene identifier."""
    idx = pd.Index([str(i).split(".")[0] if str(i).startswith("ENSMUSG") else str(i) for i in df.index])
    df = df.set_axis(idx)
    df = df[~idx.isna() & (idx != "") & (idx != "nan")]
    if not df.index.has_duplicates:
        return df
    if value_type in ("counts", "tpm", "fpkm", "cpm"):
        return df.groupby(level=0).sum()
    return df.groupby(level=0).max()


def to_log2(df: pd.DataFrame, value_type: str) -> pd.DataFrame:
    if value_type == "counts":
        lib = df.sum(axis=0)
        cpm = df.divide(lib, axis=1) * 1e6
        return np.log2(cpm + 1)
    if value_type in ("tpm", "fpkm", "cpm"):
        return np.log2(df + 1)
    return df  # vst / log2norm


def assign_group(ds: Dataset, column: str, sample: dict | None) -> str:
    if ds.group_source == "column":
        text = column
    elif ds.group_source == "title":
        text = (sample or {}).get("title", column)
    elif ds.group_source.startswith("characteristic:"):
        key = ds.group_source.split(":", 1)[1]
        text = (sample or {}).get("characteristics", {}).get(key, "")
    else:
        raise ValueError(ds.group_source)
    for pat, group in ds.groups:
        if re.search(pat, text):
            return group
    return "unassigned"


def resolve(ds: Dataset, columns: list[str]) -> dict[str, dict]:
    """column -> {gsm, title, characteristics, column_source}"""
    if not ds.resolve_gsm:
        return {c: {"gsm": "", "title": "", "characteristics": {}, "column_source": "column_labels_only"}
                for c in columns}

    smps = geo_meta.samples(ds.acc)
    if ds.gsm_filter:
        smps = [s for s in smps if re.search(ds.gsm_filter, s["title"] + " " + " ".join(s["characteristics"].values()))]

    def norm(s: str) -> str:
        return re.sub(r"[^a-z0-9]", "", re.sub(r"\([^)]*\)", "", s).lower())

    out: dict[str, dict] = {}
    used: set[str] = set()

    # explicit key regexes take precedence when configured
    if ds.column_key_regex and ds.title_key_regex:
        for c in columns:
            mc = re.search(ds.column_key_regex, c)
            if not mc:
                continue
            key = norm(mc.group(1))
            for s in smps:
                mt = re.search(ds.title_key_regex, s["title"])
                if mt and norm(mt.group(1)) == key and s["gsm"] not in used:
                    out[c] = dict(s, column_source="key_regex")
                    used.add(s["gsm"])
                    break

    for c in columns:
        if c in out:
            continue
        # 1. declared column name
        for s in smps:
            if s["gsm"] in used:
                continue
            if geo_meta.declared_column(s) == c:
                out[c] = dict(s, column_source="explicit_description")
                used.add(s["gsm"])
                break
        if c in out:
            continue
        # 2. GSM id inside the column name (per-sample files)
        for s in smps:
            if s["gsm"] in used:
                continue
            if s["gsm"] in c:
                out[c] = dict(s, column_source="supplementary_file")
                used.add(s["gsm"])
                break
        if c in out:
            continue
        # 3. normalised title / supplementary-stem match
        nc = norm(c)
        for s in smps:
            if s["gsm"] in used:
                continue
            keys = {norm(s["title"])} | {norm(re.sub(r"\.[^.]*$", "", f)) for f in s["supplementary"]}
            if nc in keys or any(k and (k.endswith(nc) or nc.endswith(k)) for k in keys):
                out[c] = dict(s, column_source="title_match")
                used.add(s["gsm"])
                break

    missing = [c for c in columns if c not in out]
    if missing and ds.gsm_filter:
        # gsm_filter restricts the sample universe; leftover columns are other
        # fractions / arms in the same matrix and are dropped, not fatal.
        return out
    if missing:
        raise RuntimeError(
            f"{ds.key}: unresolved columns {missing}; "
            f"unused samples {[s['gsm'] + '/' + s['title'] for s in smps if s['gsm'] not in used]}"
        )
    return out


def build(ds: Dataset, outdir: str) -> dict:
    path = os.path.join(RAW, ds.acc, ds.file)
    df = READERS[ds.reader](ds, path)
    df.columns = clean_columns(ds, list(df.columns))
    df = df.loc[:, ~df.columns.duplicated()]
    df = df.apply(pd.to_numeric, errors="coerce")

    mapping = resolve(ds, list(df.columns))
    df = df[[c for c in df.columns if c in mapping]]
    groups = {c: assign_group(ds, c, mapping[c]) for c in df.columns}
    keep = [c for c in df.columns if groups[c] != "unassigned"]
    dropped = [c for c in df.columns if groups[c] == "unassigned"]
    df = df[keep]

    df = collapse_genes(df, ds.value_type)
    df = df.dropna(how="all")
    expr = to_log2(df, ds.value_type)

    os.makedirs(outdir, exist_ok=True)
    expr_path = os.path.join(outdir, f"{ds.key}.expr.tsv.gz")
    expr.to_csv(expr_path, sep="\t", float_format="%.5f")

    design_path = os.path.join(outdir, f"{ds.key}.design.tsv")
    with open(design_path, "w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["column", "gsm", "geo_title", "group", "model_family", "model_class",
                    "setting", "column_source", "geo_characteristics"])
        for c in keep:
            m = mapping[c]
            chars = "; ".join(f"{k}={v}" for k, v in m.get("characteristics", {}).items())
            w.writerow([c, m.get("gsm", ""), m.get("title", ""), groups[c], ds.model_family,
                        ds.model_class, ds.setting, m.get("column_source", ""), chars])

    counts_per_group = pd.Series(groups).loc[keep].value_counts().to_dict()
    return {
        "key": ds.key,
        "accession": ds.acc,
        "file": ds.file,
        "n_genes": expr.shape[0],
        "n_samples": expr.shape[1],
        "value_type": ds.value_type,
        "gene_id_type": ds.gene_id_type,
        "groups": "; ".join(f"{k}={v}" for k, v in sorted(counts_per_group.items())),
        "dropped_columns": ",".join(dropped),
        "column_sources": ",".join(sorted({mapping[c]["column_source"] for c in keep})),
        "median_expr": f"{float(np.nanmedian(expr.to_numpy())):.3f}",
        "expr_file": os.path.basename(expr_path),
        "design_file": os.path.basename(design_path),
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--outdir", default="results/opus_gemm/data/matrices")
    ap.add_argument("--qc", default="results/opus_gemm/build_qc.tsv")
    ap.add_argument("--only", nargs="*")
    args = ap.parse_args()

    rows = []
    for ds in DATASETS:
        if args.only and ds.key not in args.only:
            continue
        try:
            row = build(ds, args.outdir)
            rows.append(row)
            print(f"[ok] {ds.key:34s} genes={row['n_genes']:6d} n={row['n_samples']:3d} "
                  f"groups: {row['groups']}", flush=True)
            if row["dropped_columns"]:
                print(f"      dropped: {row['dropped_columns']}")
        except Exception as exc:
            print(f"[FAIL] {ds.key}: {exc}", file=sys.stderr, flush=True)
            raise

    if rows:
        os.makedirs(os.path.dirname(args.qc) or ".", exist_ok=True)
        with open(args.qc, "w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]), delimiter="\t")
            w.writeheader()
            w.writerows(rows)
        print(f"\nQC -> {args.qc}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
