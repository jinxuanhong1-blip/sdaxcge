#!/usr/bin/env python3
"""NHEJ-STING logic wave.

1. Score GSE252340 (H1944, MPS1i BAY1217389 vs DMSO) as a positive control
   that a prespecified STING / type-I IFN gene set moves up when cGAS-STING
   is forcibly activated. TREX1 is part of that set (the paper's negative
   regulator).
2. On the already-cataloged public CLDN4 KD matrices (GSE207704, GSE22493),
   score the same STING set and a prespecified NHEJ effector set. These are
   not new accessions.

No new numbers are invented. Downloads are public GEO files only.
"""

from __future__ import annotations

import gzip
import io
import json
import math
import os
import urllib.request
from itertools import combinations
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from scipy import stats

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "results" / "nhej_sting_wave"
TAB = OUT / "tables"
FIG = OUT / "figures"
CACHE = Path(os.environ.get("NHEJ_CACHE", "/tmp/nhej_sting_cache"))
UA = "nhej-sting-wave/1.0 (public GEO reanalysis)"

# Locked before looking at genome-wide ranks. Symbols cover current and
# cufflinks-era names (STING1/TMEM173, CGAS/MB21D1).
STING_PANEL = [
    "TREX1",
    "STING1",
    "TMEM173",
    "CGAS",
    "MB21D1",
    "TBK1",
    "IRF3",
    "STAT1",
    "STAT2",
    "IRF7",
    "IRF9",
    "IFNB1",
    "IFIT1",
    "IFIT2",
    "IFIT3",
    "ISG15",
    "MX1",
    "MX2",
    "OAS1",
    "OAS2",
    "OAS3",
    "OASL",
    "CXCL10",
    "CCL5",
    "CXCL9",
    "CXCL11",
    "IFI27",
    "IFI44",
    "IFI44L",
    "RSAD2",
    "BST2",
    "IFIH1",
    "DDX58",
    "ISG20",
    "XAF1",
    "USP18",
    "IFITM1",
    "HLA-A",
    "HLA-B",
    "HLA-C",
    "B2M",
    "TAP1",
    "TAP2",
    "PSMB8",
    "PSMB9",
    "PSMB10",
    "NLRC5",
    "CD274",
]
# Bitler 2022 RPPA effectors of NHEJ / single-strand repair, plus the core
# NHEJ machinery. Direction expected after CLDN4 loss, if the protein claim
# is transcriptional: down.
NHEJ_PANEL = [
    "TP53BP1",
    "XRCC1",
    "LIG4",
    "XRCC4",
    "NHEJ1",
    "PRKDC",
    "DCLRE1C",
    "XRCC5",
    "XRCC6",
    "PAXX",
    "POLL",
    "POLM",
]


def fetch(url: str, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and dest.stat().st_size > 0:
        return dest
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=180) as r, dest.open("wb") as f:
        while True:
            chunk = r.read(1 << 20)
            if not chunk:
                break
            f.write(chunk)
    return dest


def bh(pvals: np.ndarray) -> np.ndarray:
    p = np.asarray(pvals, dtype=float)
    out = np.full(p.shape, np.nan)
    ok = np.isfinite(p)
    if ok.sum() == 0:
        return out
    pv = p[ok]
    n = pv.size
    order = np.argsort(pv)
    ranked = pv[order]
    q = ranked * n / (np.arange(1, n + 1))
    q = np.minimum.accumulate(q[::-1])[::-1]
    q = np.clip(q, 0, 1)
    tmp = np.empty(n)
    tmp[order] = q
    out[ok] = tmp
    return out


def load_hallmark() -> dict[str, list[str]]:
    gmt = Path(__file__).with_name("hallmark_ifn.gmt")
    sets = {}
    for line in gmt.read_text().splitlines():
        if not line.strip():
            continue
        parts = line.split("\t")
        sets[parts[0]] = parts[1:]
    return sets


def gene_set_test(logfc: pd.Series, genes: list[str], background: pd.Series) -> dict:
    """One-sided tests that the set's log2FC is shifted up."""
    present = [g for g in genes if g in logfc.index]
    # collapse alias pairs that both resolved: keep both if distinct rows
    vals = logfc.loc[present].astype(float)
    vals = vals[~vals.index.duplicated(keep="first")]
    bg = background.drop(index=vals.index, errors="ignore")
    rec = {
        "n_requested": len(genes),
        "n_present": int(vals.size),
        "n_up": int((vals > 0).sum()) if vals.size else 0,
        "n_down": int((vals < 0).sum()) if vals.size else 0,
        "median_log2fc": float(vals.median()) if vals.size else math.nan,
        "mean_log2fc": float(vals.mean()) if vals.size else math.nan,
    }
    if vals.size >= 5:
        w = stats.wilcoxon(vals.to_numpy(), alternative="greater", zero_method="wilcox")
        rec["wilcoxon_greater_p"] = float(w.pvalue)
        rec["wilcoxon_stat"] = float(w.statistic)
    else:
        rec["wilcoxon_greater_p"] = math.nan
        rec["wilcoxon_stat"] = math.nan
    if vals.size >= 5 and bg.size >= 20:
        u = stats.mannwhitneyu(vals.to_numpy(), bg.to_numpy(), alternative="greater")
        rec["mw_vs_background_p"] = float(u.pvalue)
        rec["background_n"] = int(bg.size)
        rec["background_median_log2fc"] = float(bg.median())
    else:
        rec["mw_vs_background_p"] = math.nan
        rec["background_n"] = int(bg.size)
        rec["background_median_log2fc"] = float(bg.median()) if bg.size else math.nan
    return rec


def sample_permutation_median(norm: pd.DataFrame, genes: list[str], treat_cols, ctrl_cols) -> dict:
    """Exact two-vs-two label permutations of the set median log2FC."""
    present = [g for g in genes if g in norm.index]
    sub = norm.loc[present]
    sub = sub[~sub.index.duplicated(keep="first")]
    cols = list(treat_cols) + list(ctrl_cols)
    assert len(cols) == 4

    def med(a_cols, b_cols):
        a = sub[list(a_cols)].mean(axis=1)
        b = sub[list(b_cols)].mean(axis=1)
        return float(np.median(np.log2((a + 1.0) / (b + 1.0))))

    observed = med(treat_cols, ctrl_cols)
    # all ways to choose 2 of 4 columns as the 'treatment' group
    stats_perm = []
    for chosen in combinations(cols, 2):
        other = [c for c in cols if c not in chosen]
        stats_perm.append(med(chosen, other))
    stats_perm = np.array(stats_perm, dtype=float)
    # one-sided: as large as observed
    p = float(np.mean(stats_perm >= observed - 1e-12))
    return {
        "observed_median_log2fc": observed,
        "n_partitions": int(stats_perm.size),
        "perm_p_greater": p,
        "rank_desc": int(np.sum(stats_perm > observed) + 1),
    }


def summarize_focal(logfc: pd.Series, extra: dict | None = None) -> pd.DataFrame:
    genes = []
    for g in ["CLDN4", "TREX1", "STING1", "TMEM173", "CGAS", "MB21D1"] + STING_PANEL + NHEJ_PANEL:
        if g not in genes:
            genes.append(g)
    rows = []
    for g in genes:
        if g not in logfc.index:
            rows.append({"gene": g, "present": 0, "log2fc": np.nan})
            continue
        row = {"gene": g, "present": 1, "log2fc": float(logfc.loc[g])}
        if extra and g in extra:
            row.update(extra[g])
        rows.append(row)
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# GSE252340
# ---------------------------------------------------------------------------

GSE252340_FILES = {
    "DMSO_1": "https://ftp.ncbi.nlm.nih.gov/geo/samples/GSM7999nnn/GSM7999701/suppl/GSM7999701_H1944_DMSO_1.xlsx",
    "DMSO_2": "https://ftp.ncbi.nlm.nih.gov/geo/samples/GSM7999nnn/GSM7999702/suppl/GSM7999702_H1944_DMSO_2.xlsx",
    "BAY_1": "https://ftp.ncbi.nlm.nih.gov/geo/samples/GSM7999nnn/GSM7999703/suppl/GSM7999703_H1944_BAY_1.xlsx",
    "BAY_2": "https://ftp.ncbi.nlm.nih.gov/geo/samples/GSM7999nnn/GSM7999704/suppl/GSM7999704_H1944_BAY_2.xlsx",
}
ACUTE_COL = {
    "DMSO_1": "H1944_DMSO_1_.Read_Count",
    "DMSO_2": "H1944_DMSO_2_.Read_Count",
    "BAY_1": "H1944_BAY_1_.Read_Count",
    "BAY_2": "H1944_BAY_2_.Read_Count",
}


def _sheet_by_header(path: Path, header_token: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name in wb.sheetnames:
        ws = wb[name]
        rows = ws.iter_rows(values_only=True)
        # skip a title row if the first cell is not IDs
        first = next(rows)
        if first and first[0] == "IDs":
            header = list(first)
            data = list(rows)
        else:
            header = list(next(rows))
            data = list(rows)
        if header_token in header:
            df = pd.DataFrame(data, columns=header)
            wb.close()
            return df
    wb.close()
    raise RuntimeError(f"{header_token} not in {path}")


def load_gse252340() -> tuple[pd.DataFrame, pd.DataFrame]:
    frames = []
    for key, url in GSE252340_FILES.items():
        path = fetch(url, CACHE / "gse252340" / f"{key}.xlsx")
        df = _sheet_by_header(path, ACUTE_COL[key])
        sub = df[["Gene_Symbol", ACUTE_COL[key]]].copy()
        sub.columns = ["gene", key]
        sub[key] = pd.to_numeric(sub[key], errors="coerce")
        sub = sub.dropna(subset=["gene"])
        sub["gene"] = sub["gene"].astype(str).str.strip()
        sub = sub[sub["gene"].ne("") & sub["gene"].ne("nan")]
        sub = sub.groupby("gene", as_index=False)[key].sum()
        frames.append(sub.set_index("gene"))
    acute = frames[0].join(frames[1:], how="outer").fillna(0.0)

    # Resistance sheets are duplicated in every workbook and are not in the
    # GEO series design. Read them once from the DMSO_1 file.
    path = CACHE / "gse252340" / "DMSO_1.xlsx"
    ctrl = _sheet_by_header(path, "BAY.R_DMSO_1_.Read_Count")
    treat = _sheet_by_header(path, "BAY.R_BAY_1_.Read_Count")
    c = ctrl[["Gene_Symbol", "BAY.R_DMSO_1_.Read_Count", "BAY.R_DMSO_2_.Read_Count"]].copy()
    c.columns = ["gene", "R_DMSO_1", "R_DMSO_2"]
    t = treat[["Gene_Symbol", "BAY.R_BAY_1_.Read_Count", "BAY.R_BAY_2._Read_Count"]].copy()
    t.columns = ["gene", "R_BAY_1", "R_BAY_2"]
    for df in (c, t):
        df["gene"] = df["gene"].astype(str).str.strip()
        for col in df.columns[1:]:
            df[col] = pd.to_numeric(df[col], errors="coerce")
        df.dropna(subset=["gene"], inplace=True)
    c = c.groupby("gene", as_index=False)[["R_DMSO_1", "R_DMSO_2"]].sum().set_index("gene")
    t = t.groupby("gene", as_index=False)[["R_BAY_1", "R_BAY_2"]].sum().set_index("gene")
    resistant = c.join(t, how="outer").fillna(0.0)
    return acute, resistant


def median_of_ratios(counts: pd.DataFrame) -> pd.DataFrame:
    x = counts.clip(lower=0).astype(float)
    pos = x.where(x > 0)
    log_geo = np.log(pos).mean(axis=1)
    ratios = np.exp(np.log(pos).sub(log_geo, axis=0))
    size = ratios.median(axis=0)
    size = size / size.mean()
    return x.div(size, axis=1)


def contrast_table(norm: pd.DataFrame, treat, ctrl) -> pd.DataFrame:
    a = norm[list(treat)]
    b = norm[list(ctrl)]
    log_a = np.log2(a + 1.0)
    log_b = np.log2(b + 1.0)
    mean_a = a.mean(axis=1)
    mean_b = b.mean(axis=1)
    log2fc = np.log2((mean_a + 1.0) / (mean_b + 1.0))
    # Welch on log2(norm+1). n=2; p-values are descriptive.
    tstat, p = stats.ttest_ind(log_a.to_numpy(), log_b.to_numpy(), axis=1, equal_var=False)
    out = pd.DataFrame(
        {
            "log2fc": log2fc,
            "mean_treat": mean_a,
            "mean_ctrl": mean_b,
            "welch_t": tstat,
            "welch_p": p,
        },
        index=norm.index,
    )
    out["welch_q"] = bh(out["welch_p"].to_numpy())
    return out


def score_contrast(name: str, norm: pd.DataFrame, treat, ctrl, sets: dict) -> tuple[pd.DataFrame, dict]:
    tab = contrast_table(norm, treat, ctrl)
    expressed = tab.index[tab[["mean_treat", "mean_ctrl"]].mean(axis=1) >= 10]
    bg = tab.loc[expressed, "log2fc"]
    set_rows = []
    perm = {}
    for set_name, genes in sets.items():
        rec = gene_set_test(tab["log2fc"], genes, bg)
        rec["contrast"] = name
        rec["set"] = set_name
        if norm.shape[1] == 4 and len(treat) == 2 and len(ctrl) == 2:
            perm_rec = sample_permutation_median(norm, genes, treat, ctrl)
            rec.update({f"perm_{k}": v for k, v in perm_rec.items()})
            perm[set_name] = perm_rec
        set_rows.append(rec)
    return tab, {"sets": set_rows, "perm": perm, "n_genes": int(tab.shape[0]), "n_expressed_ge10": int(expressed.size)}


# ---------------------------------------------------------------------------
# GSE207704 collapsed FPKM (already cataloged)
# ---------------------------------------------------------------------------

def load_gse207704() -> pd.DataFrame:
    path = fetch(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE207nnn/GSE207704/suppl/GSE207704_CLDN4_RNAseq.txt.gz",
        CACHE / "GSE207704_CLDN4_RNAseq.txt.gz",
    )
    df = pd.read_csv(path, sep="\t")
    cols = {
        "MCF7_CLDN4KO_FPKM (fpkm)": "MCF7_KO",
        "MCF7_WT_FPKM (fpkm)": "MCF7_WT",
        "T47D_CLDN4KO_FPKM (fpkm)": "T47D_KO",
        "T47D_WT_FPKM (fpkm)": "T47D_WT",
    }
    keep = df[["gene_short_name"] + list(cols)].copy()
    keep = keep.rename(columns=cols)
    keep["gene_short_name"] = keep["gene_short_name"].astype(str)
    for c in cols.values():
        keep[c] = pd.to_numeric(keep[c], errors="coerce")
    keep = keep.dropna(subset=["gene_short_name"])
    # duplicate symbols: sum FPKM
    return keep.groupby("gene_short_name", as_index=True)[list(cols.values())].sum()


def score_collapsed_fpkm(fpkm: pd.DataFrame, sets: dict) -> tuple[pd.DataFrame, list]:
    eps = 0.1
    lfc_mcf7 = np.log2((fpkm["MCF7_KO"] + eps) / (fpkm["MCF7_WT"] + eps))
    lfc_t47d = np.log2((fpkm["T47D_KO"] + eps) / (fpkm["T47D_WT"] + eps))
    tab = pd.DataFrame(
        {
            "log2fc_MCF7": lfc_mcf7,
            "log2fc_T47D": lfc_t47d,
            "log2fc": (lfc_mcf7 + lfc_t47d) / 2.0,
            "mean_treat": (fpkm["MCF7_KO"] + fpkm["T47D_KO"]) / 2.0,
            "mean_ctrl": (fpkm["MCF7_WT"] + fpkm["T47D_WT"]) / 2.0,
        }
    )
    expressed = tab.index[tab[["mean_treat", "mean_ctrl"]].mean(axis=1) >= 1.0]
    bg = tab.loc[expressed, "log2fc"]
    rows = []
    for set_name, genes in sets.items():
        rec = gene_set_test(tab["log2fc"], genes, bg)
        rec["contrast"] = "GSE207704_mean_of_two_lines"
        rec["set"] = set_name
        # both lines same direction
        present = [g for g in genes if g in tab.index]
        both_up = int(((tab.loc[present, "log2fc_MCF7"] > 0) & (tab.loc[present, "log2fc_T47D"] > 0)).sum())
        both_dn = int(((tab.loc[present, "log2fc_MCF7"] < 0) & (tab.loc[present, "log2fc_T47D"] < 0)).sum())
        rec["n_both_lines_up"] = both_up
        rec["n_both_lines_down"] = both_dn
        rows.append(rec)
    return tab, rows


# ---------------------------------------------------------------------------
# GSE22493 two-color log ratio, KD / OE (already cataloged)
# ---------------------------------------------------------------------------

def load_gse22493() -> pd.DataFrame:
    mat_path = fetch(
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE22nnn/GSE22493/matrix/GSE22493_series_matrix.txt.gz",
        CACHE / "GSE22493_series_matrix.txt.gz",
    )
    gpl_path = fetch(
        "https://ftp.ncbi.nlm.nih.gov/geo/platforms/GPL10nnn/GPL10555/soft/GPL10555_family.soft.gz",
        CACHE / "GPL10555_family.soft.gz",
    )
    # platform table
    symbols = {}
    with gzip.open(gpl_path, "rt", errors="replace") as f:
        in_table = False
        header = None
        for line in f:
            if line.startswith("!platform_table_begin"):
                in_table = True
                header = None
                continue
            if line.startswith("!platform_table_end"):
                break
            if not in_table:
                continue
            parts = line.rstrip("\n").split("\t")
            if header is None:
                header = parts
                continue
            rec = dict(zip(header, parts))
            pid = rec.get("ID", "").strip()
            orf = rec.get("ORF", "").strip()
            desc = rec.get("DESCRIPTION", "")
            sym = orf
            if not sym or sym in {"NA", "n/a"}:
                # "GENE--description" on this oligonucleotide platform
                sym = desc.split("--", 1)[0].strip()
            if pid and sym and sym not in {"NA", "n/a", ""}:
                symbols[pid] = sym
    # matrix
    text = gzip.open(mat_path, "rt", errors="replace").read().splitlines()
    start = text.index("!series_matrix_table_begin") + 1
    end = text.index("!series_matrix_table_end")
    df = pd.read_csv(io.StringIO("\n".join(text[start:end])), sep="\t")
    df["ID_REF"] = df["ID_REF"].astype(str).str.replace('"', "", regex=False)
    for c in df.columns[1:]:
        df[c] = pd.to_numeric(df[c].astype(str).str.replace('"', "", regex=False), errors="coerce")
    df["gene"] = df["ID_REF"].map(symbols)
    df = df.dropna(subset=["gene"])
    df["gene"] = df["gene"].astype(str)
    df = df[df["gene"].ne("") & ~df["gene"].str.startswith("LOC")]
    value_cols = [c for c in df.columns if c.startswith("GSM")]
    # mean of probes per gene (log2 KD/OE already)
    return df.groupby("gene", as_index=True)[value_cols].mean()


def score_gse22493(mat: pd.DataFrame, sets: dict) -> tuple[pd.DataFrame, list]:
    vals = mat.to_numpy(dtype=float)
    log2fc = np.nanmean(vals, axis=1)
    # one-sample t vs 0 across 3 arrays
    tstat, p = stats.ttest_1samp(vals, popmean=0.0, axis=1, nan_policy="omit")
    tab = pd.DataFrame({"log2fc": log2fc, "onesample_t": tstat, "onesample_p": p}, index=mat.index)
    tab["onesample_q"] = bh(tab["onesample_p"].to_numpy())
    bg = tab["log2fc"].dropna()
    rows = []
    for set_name, genes in sets.items():
        rec = gene_set_test(tab["log2fc"], genes, bg)
        rec["contrast"] = "GSE22493_KD_over_OE"
        rec["set"] = set_name
        rows.append(rec)
    return tab, rows


def bar_focal(tab: pd.DataFrame, path: Path, title: str) -> None:
    genes = [
        "TREX1",
        "STING1",
        "CGAS",
        "TBK1",
        "IRF3",
        "STAT1",
        "IFNB1",
        "IFIT1",
        "ISG15",
        "MX1",
        "OAS1",
        "CXCL10",
        "CCL5",
        "TAP1",
        "B2M",
        "CD274",
        "TP53BP1",
        "XRCC1",
        "LIG4",
        "PRKDC",
    ]
    sub = tab.reindex(genes).dropna(subset=["log2fc"])
    fig, ax = plt.subplots(figsize=(10.2, 4.2))
    colors = ["#b45309" if g in ("TP53BP1", "XRCC1", "LIG4", "PRKDC") else "#1d4e89" for g in sub.index]
    ax.bar(np.arange(len(sub)), sub["log2fc"], color=colors)
    ax.axhline(0, color="#444", lw=0.8)
    ax.set_xticks(np.arange(len(sub)))
    ax.set_xticklabels(sub.index, rotation=60, ha="right")
    ax.set_ylabel("log2FC (BAY1217389 / DMSO)")
    ax.set_title(title)
    fig.tight_layout()
    fig.savefig(path, dpi=140)
    plt.close(fig)


def main() -> None:
    TAB.mkdir(parents=True, exist_ok=True)
    FIG.mkdir(parents=True, exist_ok=True)
    hallmark = load_hallmark()
    sets = {
        "STING_IFN_panel": STING_PANEL,
        "NHEJ_panel": NHEJ_PANEL,
        "Hallmark_Interferon_Alpha": hallmark["Interferon_Alpha_Response"],
        "Hallmark_Interferon_Gamma": hallmark["Interferon_Gamma_Response"],
    }

    acute, resistant = load_gse252340()
    acute_norm = median_of_ratios(acute)
    res_norm = median_of_ratios(resistant)
    acute_tab, acute_meta = score_contrast(
        "GSE252340_acute_BAY_vs_DMSO",
        acute_norm,
        ["BAY_1", "BAY_2"],
        ["DMSO_1", "DMSO_2"],
        sets,
    )
    res_tab, res_meta = score_contrast(
        "GSE252340_workbook_BAY.R_vs_DMSO",
        res_norm,
        ["R_BAY_1", "R_BAY_2"],
        ["R_DMSO_1", "R_DMSO_2"],
        sets,
    )
    acute_tab.to_csv(TAB / "gse252340_acute_gene_log2fc.tsv", sep="\t", float_format="%.6g")
    focal = acute_tab.loc[acute_tab.index.intersection(STING_PANEL + NHEJ_PANEL + ["CLDN4"]), ["log2fc", "welch_p", "welch_q", "mean_treat", "mean_ctrl"]]
    focal.to_csv(TAB / "gse252340_acute_focal.tsv", sep="\t", float_format="%.6g")
    res_focal = res_tab.loc[res_tab.index.intersection(STING_PANEL + NHEJ_PANEL + ["CLDN4"]), ["log2fc", "welch_p", "welch_q"]]
    res_focal.to_csv(TAB / "gse252340_resistant_sheet_focal.tsv", sep="\t", float_format="%.6g")
    bar_focal(
        acute_tab,
        FIG / "gse252340_acute_focal_log2fc.png",
        "GSE252340 H1944 acute MPS1i (n=2 vs 2): STING/IFN genes in blue, NHEJ genes in brown",
    )

    fpkm = load_gse207704()
    t47_tab, t47_rows = score_collapsed_fpkm(fpkm, sets)
    t47_focal = t47_tab.loc[
        t47_tab.index.intersection(STING_PANEL + NHEJ_PANEL + ["CLDN4"]),
        ["log2fc", "log2fc_MCF7", "log2fc_T47D", "mean_treat", "mean_ctrl"],
    ]
    t47_focal.to_csv(TAB / "gse207704_nhej_sting_focal.tsv", sep="\t", float_format="%.6g")

    arr = load_gse22493()
    arr_tab, arr_rows = score_gse22493(arr, sets)
    arr_focal = arr_tab.loc[arr_tab.index.intersection(STING_PANEL + NHEJ_PANEL + ["CLDN4"])]
    arr_focal.to_csv(TAB / "gse22493_nhej_sting_focal.tsv", sep="\t", float_format="%.6g")

    set_df = pd.DataFrame(acute_meta["sets"] + res_meta["sets"] + t47_rows + arr_rows)
    # stable column order
    front = ["contrast", "set", "n_present", "n_up", "n_down", "median_log2fc", "mean_log2fc", "wilcoxon_greater_p", "mw_vs_background_p"]
    cols = front + [c for c in set_df.columns if c not in front]
    set_df = set_df[cols]
    set_df.to_csv(TAB / "geneset_tests.tsv", sep="\t", index=False, float_format="%.6g")

    summary = {
        "gse252340_acute": {
            "n_genes": acute_meta["n_genes"],
            "n_expressed_ge10": acute_meta["n_expressed_ge10"],
            "library_sizes_raw": {c: float(acute[c].sum()) for c in acute.columns},
            "sets": acute_meta["sets"],
            "CLDN4_log2fc": None if "CLDN4" not in acute_tab.index else float(acute_tab.loc["CLDN4", "log2fc"]),
            "TREX1_log2fc": None if "TREX1" not in acute_tab.index else float(acute_tab.loc["TREX1", "log2fc"]),
            "n_welch_q_lt_0.05": int((acute_tab["welch_q"] < 0.05).sum()),
        },
        "gse252340_resistant_sheet": {
            "note": "Sheets labeled BAY.R are duplicated inside every sample workbook and are not described in the GEO series summary.",
            "sets": res_meta["sets"],
            "TREX1_log2fc": None if "TREX1" not in res_tab.index else float(res_tab.loc["TREX1", "log2fc"]),
        },
        "gse207704": {"sets": t47_rows, "CLDN4": None if "CLDN4" not in t47_tab.index else {
            "log2fc_mean": float(t47_tab.loc["CLDN4", "log2fc"]),
            "MCF7": float(t47_tab.loc["CLDN4", "log2fc_MCF7"]),
            "T47D": float(t47_tab.loc["CLDN4", "log2fc_T47D"]),
        }},
        "gse22493": {"sets": arr_rows, "CLDN4": None if "CLDN4" not in arr_tab.index else {
            "log2fc": float(arr_tab.loc["CLDN4", "log2fc"]),
            "p": float(arr_tab.loc["CLDN4", "onesample_p"]),
        }},
    }
    (TAB / "summary.json").write_text(json.dumps(summary, indent=2))
    print(set_df.to_string(index=False))
    print("--- focal acute TREX1/STING/NHEJ ---")
    show = ["TREX1", "STING1", "CGAS", "IFNB1", "IFIT1", "ISG15", "CXCL10", "CCL5", "TAP1", "TP53BP1", "XRCC1", "CLDN4"]
    print(acute_tab.reindex(show)[["log2fc", "welch_p", "welch_q"]].to_string())
    print("--- 207704 CLDN4 / 53BP1 / XRCC1 / TREX1 ---")
    print(t47_tab.reindex(["CLDN4", "TP53BP1", "XRCC1", "TREX1", "STING1", "TMEM173", "ISG15", "IFIT1"])[
        ["log2fc", "log2fc_MCF7", "log2fc_T47D"]
    ].to_string())
    print("--- 22493 ---")
    print(arr_tab.reindex(["CLDN4", "TP53BP1", "XRCC1", "TREX1", "TMEM173", "STING1", "ISG15", "IFIT1", "CXCL10"])[
        ["log2fc", "onesample_p", "onesample_q"]
    ].to_string())


if __name__ == "__main__":
    main()
