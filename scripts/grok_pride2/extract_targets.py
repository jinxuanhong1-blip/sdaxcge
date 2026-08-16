#!/usr/bin/env python3
"""Scan downloaded protein tables for PD-1 / PD-L1 / TROP2 / CLDN4 rows."""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "results" / "grok_pride2"
TABLES = RESULTS / "protein_tables"

TARGETS = {
    "PDCD1": re.compile(r"\b(PDCD1|PD-?1|CD279)\b", re.I),
    "CD274": re.compile(r"\b(CD274|PD-?L1|B7-?H1|PDCD1LG1)\b", re.I),
    "PDCD1LG2": re.compile(r"\b(PDCD1LG2|PD-?L2|CD273|B7-?DC)\b", re.I),
    "TACSTD2": re.compile(r"\b(TACSTD2|TROP-?2|EGP-?1|GA733-?1)\b", re.I),
    "CLDN4": re.compile(r"\b(CLDN4|claudin[- ]?4)\b", re.I),
}


def iter_rows(path: Path):
    if path.suffix.lower() in {".xlsx", ".xls"}:
        try:
            import openpyxl  # type: ignore
        except Exception:
            return
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                yield ["" if v is None else str(v) for v in row]
        return
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    # sniff delimiter
    sample = text[:4000]
    delim = "\t" if sample.count("\t") >= sample.count(",") else ","
    reader = csv.reader(text.splitlines(), delimiter=delim)
    for row in reader:
        yield row


def main() -> None:
    hits = []
    scanned = []
    if not TABLES.exists():
        raise SystemExit("no protein_tables dir")
    for path in sorted(TABLES.rglob("*")):
        if not path.is_file():
            continue
        if path.suffix.lower() in {".zip", ".gz", ".tar"}:
            scanned.append({"file": str(path.relative_to(ROOT)), "status": "archive_not_expanded"})
            continue
        if path.stat().st_size > 60 * 1024 * 1024:
            scanned.append({"file": str(path.relative_to(ROOT)), "status": "too_large"})
            continue
        n = 0
        file_hits = 0
        try:
            for i, row in enumerate(iter_rows(path)):
                n += 1
                line = " ".join(row)
                matched = [name for name, rx in TARGETS.items() if rx.search(line)]
                if matched:
                    file_hits += 1
                    hits.append(
                        {
                            "file": str(path.relative_to(ROOT)),
                            "accession": path.parent.name,
                            "row_index": i,
                            "targets": ";".join(matched),
                            "row": line[:500],
                        }
                    )
                if n > 200000:
                    break
        except Exception as exc:  # noqa: BLE001
            scanned.append({"file": str(path.relative_to(ROOT)), "status": f"error:{exc}"})
            continue
        scanned.append(
            {
                "file": str(path.relative_to(ROOT)),
                "status": "ok",
                "n_rows": n,
                "n_target_rows": file_hits,
            }
        )

    out = RESULTS / "target_protein_hits.csv"
    with out.open("w", newline="", encoding="utf-8") as fh:
        fields = ["accession", "file", "row_index", "targets", "row"]
        w = csv.DictWriter(fh, fieldnames=fields)
        w.writeheader()
        w.writerows(hits)
    (RESULTS / "target_scan_log.json").write_text(
        json.dumps({"n_files": len(scanned), "n_hits": len(hits), "files": scanned}, indent=2),
        encoding="utf-8",
    )
    print(f"scanned {len(scanned)} files; {len(hits)} target rows")


if __name__ == "__main__":
    main()
