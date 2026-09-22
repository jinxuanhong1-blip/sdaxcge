#!/usr/bin/env python3
"""Compile the master story-fit PPT table from numbers already published in PRs.

No new statistical fits. Every coefficient is a string copied from an existing
PR body or from the public handoff. Do not edit a number here without a PR.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent

# fit is the stricter of the source calls when two PRs disagree.
# largest_effect is the largest eligible story-direction effect the source PR
# itself reported. locked_primary is the pre-specified number when it differs.
ROWS = [
    # ----- Part 1A. Tacstd2 / TROP2 immune-cold -----
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C1",
        limb="CosMx He2022: TACSTD2-high tumor vs CD8+NK neighbors",
        fit="YES",
        largest_effect="CD8+NK count hi/lo 0.455 at 10 µm and 0.670 at 20 µm; 8/8 sections and 5/5 patients; sign P=0.031. Immune-fraction ratios 0.519 / 0.643. CD8-only count ratios 0.502 / 0.726.",
        locked_primary="Same ratios. This is the short-range TACSTD2 result, not a replacement for locked CLDN4 0.36 / 0.52.",
        n_unit="295,877 tumor cells; 8 sections; 5 patients",
        source_pr="#726 (table in #735, #738, #739)",
        caveat="At 50 µm the count ratio is 0.854 (7/8); at 100 µm 0.947 (4/8). Inside CLDN4-low the 10 µm immune-fraction ratio stays 0.507 (8/8 and 5/5); rank attenuation after CLDN4 is only 0.124. Not accounted for by the cell's own CLDN4.",
        slide_large="0.455× at 10 µm; 0.670× at 20 µm; 8/8 and 5/5; sign P=0.031",
        slide_lock="Same. Not a replacement for CLDN4 0.36 / 0.52.",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C2",
        limb="OncoSG LUAD: TACSTD2 vs T/NK-class scores after published PURITY",
        fit="YES",
        largest_effect="IMSIG NK partial ρ=−0.421; IMSIG T partial ρ=−0.404 (n=169).",
        locked_primary="CD8A partial ρ=−0.309 (95% CI −0.440 to −0.165, p=4.69e-5). GEP18 partial ρ=−0.349. A1 immune partial ρ=−0.318. Unadjusted CD8A ρ=−0.380.",
        n_unit="n=169 (portal lists 181; public z-score matrix has 169)",
        source_pr="#139 (restated #735, #736, #739)",
        caveat="East-Asian surgical LUAD, not ICI. Do not quote IMSIG neutrophil partial ρ=−0.456 as CD8.",
        slide_large="IMSIG NK partial ρ=−0.421; IMSIG T −0.404",
        slide_lock="CD8A partial ρ=−0.309 (p=4.69e-5); GEP18 −0.349",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C3",
        limb="TCGA LUAD+LUSC: TACSTD2 vs CD8A after ABSOLUTE purity",
        fit="YES",
        largest_effect="LUSC partial ρ=−0.244 (p=4.3e-8, n=493).",
        locked_primary="Pooled partial ρ=−0.191 (p=1.31e-9, n=995). LUAD partial ρ=−0.104 (p=0.020, n=502). CYT pooled partial ρ=−0.151.",
        n_unit="995 tumors with ABSOLUTE purity",
        source_pr="#107 (restated #739)",
        caveat="LUSC is stronger than LUAD (Fisher p=0.023). Single-gene CD8A, not a deconvolution fraction. Not ICI.",
        slide_large="LUSC partial ρ=−0.244 (n=493)",
        slide_lock="Pooled ρ=−0.191 (n=995); LUAD −0.104",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C4",
        limb="TCGA ESTIMATE ImmuneScore after ABSOLUTE purity",
        fit="PARTIAL",
        largest_effect="LUSC partial ρ=−0.131 (p=0.003, n=493). GEP18 is also LUSC-only (partial ρ=−0.226).",
        locked_primary="Pooled ImmuneScore partial ρ=−0.107 (p=7.4e-4, n=995). LUAD partial ρ=+0.072 (p=0.109).",
        n_unit="995 tumors",
        source_pr="#107",
        caveat="LUSC-only. Do not quote the pooled ImmuneScore as a LUAD finding.",
        slide_large="LUSC ImmuneScore partial ρ=−0.131",
        slide_lock="Pooled −0.107; LUAD +0.072 (not significant)",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C5",
        limb="TCGA eight primary cohorts: TACSTD2 vs CD8 after keratin",
        fit="PARTIAL",
        largest_effect="All-patient KRT5/6 sweep, CD3/CD8 z-mean: pooled ρ=−0.119 (95% CI −0.176 to −0.062, p=4.98e-5, 7/8 negative).",
        locked_primary="Pre-specified KRT8+KRT18+KRT19, CD8A+CD8B mean: pooled ρ=−0.069 (95% CI −0.135 to −0.002, p=0.043, I²=76%, 6/8).",
        n_unit="LUAD 516, LUSC 501, BRCA 1095, CESC 304, KIRC 533, STAD 412, BLCA 406, PAAD 178",
        source_pr="#593 primary; #705 largest all-patient sweep",
        caveat="The sweep maximum inflates |ρ|. CESC partial ρ=+0.118 on the pre-specified keratin control. Bulk association is not spatial exclusion. KRT5/6-low tertile grid max ρ=−0.152 is a subset, not the all-patient row.",
        slide_large="Sweep ρ=−0.119 (7/8; p=4.98e-5)",
        slide_lock="KRT8/18/19 ρ=−0.069 (6/8; p=0.043; I²=76%)",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C6",
        limb="CPTAC LUAD: TROP2 protein vs xCell CD8 / immune",
        fit="PARTIAL",
        largest_effect="LUAD TROP2 protein vs xCell immune ρ=−0.309 (n=110, p=0.0010); vs xCell CD8 ρ=−0.289 (p=0.0022).",
        locked_primary="After WES purity, CD8 partial ρ=−0.261 (n=108, p=0.0067); immune partial ρ=−0.272.",
        n_unit="LUAD n=110; LSCC n=108 does not replicate",
        source_pr="#99 (restated #735, #739)",
        caveat="LUAD-only. LSCC CD8 ρ=−0.080 (p=0.41). Treatment-naive surgery. RNA TACSTD2 vs xCell is null (|ρ|≤0.12). TROP2 vs CD8A protein CIs cross 0 in both histologies (#721).",
        slide_large="LUAD xCell immune ρ=−0.309; CD8 −0.289",
        slide_lock="WES-partial CD8 ρ=−0.261. LSCC null.",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C7",
        limb="GSE31210 primary LUAD: TACSTD2 vs CD8A / ESTIMATE ImmuneScore",
        fit="PARTIAL",
        largest_effect="Unadjusted ImmuneScore ρ=−0.306 (p=2.75e-6); unadjusted CD8A ρ=−0.289 (p=9.76e-6).",
        locked_primary="After ESTIMATE, partial ρ=−0.023 (ImmuneScore, p=0.735) and −0.107 (CD8A, p=0.108).",
        n_unit="n=226 primary tumors",
        source_pr="#736",
        caveat="The negative sign does not survive the ESTIMATE partial. East-Asian surgical LUAD.",
        slide_large="Unadj ImmuneScore ρ=−0.306; CD8A −0.289",
        slide_lock="ESTIMATE partial −0.023 / −0.107 (not significant)",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C8",
        limb="Public mouse: Tacstd2 % epithelial vs T fraction",
        fit="PARTIAL",
        largest_effect="GSE295824 Spearman ρ=−0.612 (n=16, p=0.012). Partial given Cldn4 ρ=−0.468 (p=0.082).",
        locked_primary="Five-study pool, within-study rank z: total c=−0.277, permutation p=0.137, n=34. Indirect path through Cldn4 includes 0.",
        n_unit="16 mice in the one negative study; pool n=34, k=5",
        source_pr="#724 (study call in #739)",
        caveat="Not a KL-vs-KP experiment. GSE264739 is opposite (ρ=+0.829, p=0.058). Do not merge with private 8KL. The n=4 ceiling ρ=−1.000 (exact p=0.083) is exploratory (#743), not this row.",
        slide_large="GSE295824 ρ=−0.612 (n=16, p=0.012)",
        slide_lock="Five-study pool c=−0.277, p=0.137 (n=34)",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C9",
        limb="GSE137244 KL vs KP: IFN-compact score",
        fit="PARTIAL",
        largest_effect="IFN compact (15 genes) Δ=−0.871, exact MW p=0.0317. Libraries overlap.",
        locked_primary="Same. APM MHC-I Δ=−0.500, p=0.31 (not significant). STING core is higher in KL (Δ=+0.475), not lower.",
        n_unit="5 KL vs 5 KP libraries",
        source_pr="#685",
        caveat="The Tacstd2 split is completely genotype-confounded (every KL library is above every KP library). A 192-cell sweep found 0 cells with Tacstd2 and Cldn4 up together with NHEJ down and IFN up.",
        slide_large="IFN-compact Δ=−0.871, MW p=0.0317",
        slide_lock="Same. APM Δ=−0.500, p=0.31 (ns).",
    ),
    dict(
        part="Part 1 — Tacstd2 cold",
        limb_id="P1-C10",
        limb="GSE131907: TACSTD2 %pos is higher in malignant / epithelial cells than in T/NK",
        fit="YES",
        largest_effect="MPE carcinoma-like TACSTD2 %pos 83.8 vs T 1.4 / NK 3.1 / B 1.3 / myeloid 5.6.",
        locked_primary="Tumor-site epithelial TACSTD2 %pos 75.0 vs T 2.8 / NK 3.2 / B 2.6 / myeloid 9.7 (#230).",
        n_unit="MPE carcinoma-like cells n=259; tumor-site comparison is cell-compartment %pos, not a patient ρ",
        source_pr="#745; tumor-site gap #230",
        caveat="This is a dissociated detection gap, not patient-level exclusion. MPE sample Spearman is underpowered (3 PE samples with ≥20 carcinoma-like cells). Tumor-site TACSTD2 vs CD8 ρ=+0.11 (n=36, p=0.54, #230).",
        slide_large="MPE TACSTD2 %pos 83.8 vs T 1.4 / NK 3.1",
        slide_lock="Tumor-site epithelial 75.0 vs T 2.8 / NK 3.2",
    ),
    # ----- Part 1B. Tacstd2 tracks TJ / CLDN4 -----
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T1",
        limb="GSE137244 KL vs KP: Tacstd2, Cldn4, and TJ scores",
        fit="YES",
        largest_effect="Cldn4 Δ=+5.570. Cldn4-edge (11 genes) Δ=+2.449. TJ_TISMO (7 genes) Δ=+3.269. Tacstd2 Δ=+3.238. TJ epithelial (18) Δ=+1.627.",
        locked_primary="Tacstd2 +3.238, Cldn4 +5.570, and every TJ score above except the 107-gene within-arm neighborhood, all with exact MW p=0.00794 and complete KL>KP separation.",
        n_unit="5 vs 5 libraries",
        source_pr="#685 (deltas restated #746, #743, #724)",
        caveat="Handoff TJ +3.03 is a different average. Do not relabel TJ_TISMO +3.269 as +3.03. In vivo Tacstd2, not TJ: GSE6135 Δ=+1.795 (7 vs 5, p=0.00253); GSE164758 Δ=+0.858 (9 vs 8, p=8.2e-5) (#743).",
        slide_large="Cldn4 Δ=+5.570; TJ_TISMO Δ=+3.269; Tacstd2 Δ=+3.238",
        slide_lock="Complete separation, exact MW p=0.00794 (5 vs 5)",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T2",
        limb="Concordant-4 malignant TACSTD2 Q4 vs Q1: junction and keratin up",
        fit="YES",
        largest_effect="ORA: Hallmark apical junction rank 1, enrichment 5.89, FDR 1.3e-5. Keratinization rank 3, enrichment 13.98, same FDR.",
        locked_primary="GSEA keratinization NES=+2.83 (rank 2 of positive NES, FDR 0.001). KEGG tight junction NES=+2.04 (FDR 0.001) but NES-rank 22. TJ family logFC=+0.223, FDR 0.0042 (206 genes, TACSTD2 held out).",
        n_unit="DE contrast 19 vs 15; expression units 64 (not n=65)",
        source_pr="#741",
        caveat="Broad TJ vs T/NK ρ=+0.137 (p=0.60); TJ-core vs T/NK ρ=−0.191 (p=0.41). Do not write that TJ-high excludes T/NK. IFN and MHC are not down on this TACSTD2 split. Locked CLDN4 %pos vs T/NK stays ρ=−0.531.",
        slide_large="ORA apical junction rank 1 (enrichment 5.89, FDR 1.3e-5)",
        slide_lock="Keratin NES=+2.83 (rank 2); KEGG TJ NES=+2.04 (rank 22)",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T3",
        limb="GSE31210: TACSTD2-high GSEA, tight junction and keratin up, EMT down",
        fit="YES",
        largest_effect="KEGG tight junction NES=+2.00 (primary FDR 0.00183). Keratinization NES=+1.91. Hallmark EMT NES=−2.15.",
        locked_primary="Same three calls are the pre-specified A8 pass. Focal CLDN4 MAS5 Δ=+704, FDR=2e-9.",
        n_unit="n=226",
        source_pr="#736",
        caveat="Hallmark apical junction NES=−1.16 (FDR 0.099), so not every junction set is up. Surgical LUAD, not ICI.",
        slide_large="KEGG TJ NES=+2.00; keratin +1.91; EMT −2.15",
        slide_lock="Same. Focal CLDN4 MAS5 Δ=+704, FDR=2e-9",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T4",
        limb="OncoSG: TACSTD2-high keratin up and EMT down; KEGG TJ not up",
        fit="PARTIAL",
        largest_effect="Keratinization NES=+1.92 (FDR 0.002). Hallmark EMT NES=−2.33 (FDR 0.002).",
        locked_primary="KEGG tight junction NES=−1.00 (FDR 0.32). Hallmark apical junction NES=−1.48 (FDR 0.0045), down rather than up. Focal CLDN4 z Δ=+1.28, FDR=2e-6.",
        n_unit="n=169",
        source_pr="#736",
        caveat="Keratin and EMT match the slide. The KEGG tight-junction set does not.",
        slide_large="Keratin NES=+1.92; EMT NES=−2.33",
        slide_lock="KEGG TJ NES=−1.00 (not up). CLDN4 z Δ=+1.28",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T5",
        limb="TCGA histology-split: TACSTD2-high tight-junction GSEA",
        fit="YES",
        largest_effect="LUSC tight-junction assembly NES=2.18, FDR=0.004.",
        locked_primary="LUAD KEGG tight junction NES=2.09, FDR=0.005. Both are within-histology, not a squamous-mix artifact.",
        n_unit="TCGA-LUAD and TCGA-LUSC, TACSTD2-high vs low",
        source_pr="#200",
        caveat="EMT-down does not replicate as pan-NSCLC (pooled NES +1.30, FDR 0.50). Strict triple intersection (TCGA NSCLC ∩ GSE207422 ∩ GSE131907, r≥0.20, FDR<0.05) is 20 genes and includes CLDN4 and ELF3.",
        slide_large="LUSC TJ assembly NES=2.18 (FDR 0.004)",
        slide_lock="LUAD KEGG TJ NES=2.09 (FDR 0.005)",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T6",
        limb="GSE131907 tumor-site epithelium: TACSTD2 prerank tight-junction GSEA",
        fit="YES",
        largest_effect="CUSTOM_TJ_CORE NES=2.21, FDR=0.",
        locked_primary="Still NES=2.17, FDR=0, after removing CLDN4 and TACSTD2. KEGG TJ NES=1.74, FDR=0.0075.",
        n_unit="n=36 tumor-site samples",
        source_pr="#230",
        caveat="Keratinization NES=1.48 (FDR 0.031) is the weaker companion. Dissociated 10x, treatment-naive, not spatial.",
        slide_large="TJ-core NES=2.21 (FDR 0)",
        slide_lock="NES=2.17 after dropping CLDN4 and TACSTD2",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T7",
        limb="GSE137244 and GSE165641: is TJ the top Tacstd2-high pathway?",
        fit="PARTIAL",
        largest_effect="Focused epithelial TJ-core NES=2.34, rank 1 on GSE137244; GSE165641 KL GEMM NES=2.17, rank 1 (#746).",
        locked_primary="Broad 589-set universe (#744): no TJ, adhesion, or claudin-family set is in ranks 1–3. GSE137244 best TJ rank is 21 (GOBP TJ organization NES=+1.99, FDR 0.086).",
        n_unit="GSE137244 is 5 vs 5 and genotype-confounded",
        source_pr="#746 focused rank; #744 broad rank",
        caveat="Do not write “TJ is pathway #1” from the broad ranking. Closest broad adhesion call is KEGG CAMs rank 10 in GSE164758 (NES=+2.06).",
        slide_large="Focused TJ-core NES=2.34, rank 1 (#746)",
        slide_lock="Broad universe: best TJ rank 21, not top 3 (#744)",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T8",
        limb="Public mouse epithelium: Tacstd2-high vs low TJ_TISMO module",
        fit="YES",
        largest_effect="7/7 mice Tacstd2-high > low; mean Δ=+0.261; binomial p=0.0078; all 7 also MW p<0.05. Rank enrichment p=3.6e-9.",
        locked_primary="Same frozen TJ_TISMO module.",
        n_unit="7 mice",
        source_pr="#743",
        caveat="Public integrate cohorts only. Not merged with private 8KL.",
        slide_large="7/7 mice, mean Δ=+0.261, binomial p=0.0078",
        slide_lock="Same frozen TJ_TISMO module",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T9",
        limb="GSE131907 MPE: TACSTD2 Q4 vs Q1 TJ-core",
        fit="YES",
        largest_effect="TJ-core Δ=+0.16, p=9.5e-8.",
        locked_primary="Holds after dropping CLDN4 (p=4.7e-8). CLDN4 co-detection OR=5.00 (88% vs 60%). Primary tS sample ρ(TACSTD2, CLDN4)=0.77 (n=10, p=0.0092).",
        n_unit="MPE carcinoma-like cells (author Malignant label is 0 in PE)",
        source_pr="#745",
        caveat="Does not replace locked author-malignant CLDN4–T/NK results. Sample-level TACSTD2 vs T/NK is not supported (3 PE samples).",
        slide_large="TJ-core Δ=+0.16, p=9.5e-8",
        slide_lock="Still p=4.7e-8 after dropping CLDN4",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T10",
        limb="DepMap / Gygi CCLE lung lines: TROP2 with CLDN4 and TJ scores",
        fit="YES",
        largest_effect="RNA NSCLC, CLDN4-edge (11 genes) ρ=+0.696 (n=143, p=4.92e-22). Protein grid maximum ρ=0.846 (NSCLC primary lines, partial on subtype, n=17, p=7.1e-5).",
        locked_primary="Gygi S1 lung TenPx TACSTD2–CLDN4 protein ρ=0.693 (n=45, p=1.31e-7). RNA CLDN4 ρ=+0.667; TJ epithelial ρ=+0.668; TJ TISMO ρ=+0.663 (all n=143).",
        n_unit="RNA n=143 NSCLC lines; protein n=45 lung lines",
        source_pr="#740 RNA and n=45 protein; #697 protein grid maximum",
        caveat="The 0.846 interval still covers 0.69. Do not quote an RPPA n≈118 (no TROP2/CLDN4 antibodies). Hallmark IFN-γ ρ=+0.339 is positive, so lines are not immune-cold. No T cells and no ICI labels.",
        slide_large="RNA CLDN4-edge ρ=+0.696 (n=143); protein max ρ=0.846 (n=17)",
        slide_lock="Protein ρ=0.693 (n=45). RNA CLDN4 ρ=+0.667",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T11",
        limb="TCGA: TACSTD2 tracks CLDN4 after KRT8/18/19",
        fit="YES",
        largest_effect="TACSTD2–CLDN4 keratin-partial pooled ρ=0.315 (19/21 carcinomas, p=3.82e-20).",
        locked_primary="Same. TACSTD2–CLDN7 pooled ρ=0.242 (19/21, p=4.35e-13). LUAD/LUSC plus the keratin funnel are 8/8 for both.",
        n_unit="21 epithelial carcinomas, primary tumors",
        source_pr="#593",
        caveat="COAD and READ are the two non-positive cohorts. COXPRESdb: mouse Tacstd2→Cldn4 is rank 1; human TACSTD2→CLDN4 is rank 3 (#200).",
        slide_large="Pooled ρ=0.315, 19/21, p=3.82e-20",
        slide_lock="CLDN7 ρ=0.242 (19/21). Mouse Tacstd2→Cldn4 rank 1 (#200)",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T12",
        limb="GSE131907: CLDN4-positive cells co-detect TACSTD2",
        fit="YES",
        largest_effect="Primary tS cells, CLDN4+ vs CLDN4−: TACSTD2 detection 92.0% vs 60.7%, OR=7.45.",
        locked_primary="Sample-mean Spearman ρ=0.770 (n=10 primary tS samples, p=0.0092). MPE carcinoma-like OR=5.00 (88.4% vs 60.5%, p=3.7e-5).",
        n_unit="6,352 primary tS cells; 10 samples; MPE carcinoma-like n=259",
        source_pr="#635 (MPE OR also in #745)",
        caveat="MPE double-positive expression levels are null (ρ=−0.025). ELF3 co-detection OR=22.6 is a different gene.",
        slide_large="Primary tS OR=7.45 (92.0% vs 60.7%)",
        slide_lock="Sample ρ=0.770 (n=10). MPE OR=5.00",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T13",
        limb="TCGA-LUAD: structural 15-gene TJ vs CD8 / GEP",
        fit="YES",
        largest_effect="TJ-15 vs CD8 Spearman ρ=−0.29 (p=1.6e-11); vs GEP ρ=−0.28 (p=9.0e-11).",
        locked_primary="ESTIMATE-purity partial ρ=−0.26 (p=1.7e-9, n=511).",
        n_unit="n=515 tumors (partial n=511)",
        source_pr="#90",
        caveat="Broad KEGG tight junction vs CD8 is ρ≈+0.07 (p≈0.10) and is not this limb. This is bulk RNA, not scRNA exclusion.",
        slide_large="TJ-15 vs CD8 ρ=−0.29 (p=1.6e-11)",
        slide_lock="Purity partial ρ=−0.26 (n=511)",
    ),
    dict(
        part="Part 1 — Tacstd2 → TJ",
        limb_id="P1-T14",
        limb="CosMx: epithelial program that contains CLDN4 and TACSTD2 vs CD8 at 50 µm",
        fit="YES",
        largest_effect="Hotspot-smoothed 10-gene program, outer quartiles, patient-mean Spearman ρ=−0.228; 5/5 patients and 8/8 sections; toroidal-shift p=0.005. High/low CD8-neighbor count ≈0.61×.",
        locked_primary="Keratin-free adhesion score (CLDN4, CDH1, EPCAM, TACSTD2) ρ=−0.213, still 5/5 and 8/8. Keratin-residualized program ρ=−0.081.",
        n_unit="225 FOVs",
        source_pr="#643 (restated #738)",
        caveat="Continuous CLDN4 alone is ρ=−0.065 (4/5; Lung6 positive). The spec is the most negative 5/5 result in a 2,754-score grid. Only CLDN4 is on the 960-plex. Does not replace locked 0.36 / 0.52.",
        slide_large="Program vs CD8 ρ=−0.228; count ≈0.61×; 5/5 and 8/8",
        slide_lock="Adhesion-only ρ=−0.213. CLDN4 alone ρ=−0.065",
    ),
    # ----- Part 2. CLDN4 pin -----
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-1",
        limb="Concordant-4: malignant CLDN4 vs patient T/NK fraction",
        fit="YES",
        largest_effect="scVI CLDN4-high fraction vs T/NK-among-malignant, within-cohort rank: DL ρ=−0.736 (I²=22%); LR χ²=36.25. Cohorts −0.764 / −0.517 / −0.802 / −0.867. I²=0 neighbor: χ²=41.03, ρ=−0.702 (#703).",
        locked_primary="%pos (UMI>0) DL ρ=−0.531 (95% CI −0.697 to −0.312, p=1.65e-5, I²=0%, N=65). Stacked Q4 vs Q1 Cliff δ=−0.724 (19/16, p=2.88e-4). Joint maximum of 6,435 count panels on both |ρ| and |Cliff δ| (#712). I²=0 partial on KRT18 is ρ=−0.533.",
        n_unit="N=65 (13+21+22+9). Not cell counts.",
        source_pr="#539 / #503 locked; #712 joint max; #703 largest |ρ|",
        caveat="The scVI |ρ| is a searched maximum (permutation p≤0.0015). Sign-only ρ=−0.552 has I²=67% and is not promoted. Largest single locked cohort is GSE123902 ρ=−0.659. Patient-bootstrap median Δ of T/NK = −0.339 (interval −0.444 to −0.093, permutation p=0.003, #714). Do not add GSE148071, GSE127465, GSE154826, GSE200563, or E-MTAB-13526.",
        slide_large="scVI fraction ρ=−0.736 (I²=22%; χ²=36.25)",
        slide_lock="%pos ρ=−0.531 (I²=0%, N=65); Cliff δ=−0.724",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-2",
        limb="GSE131907 author-malignant CLDN4 vs patient T/NK",
        fit="YES",
        largest_effect="Mean log-normalized CLDN4 ρ=−0.532 (p=0.013, n=21).",
        locked_primary="%pos ρ=−0.478 (p=0.028). Q4 vs Q1 rank-biserial r=−0.600 (6 vs 5, p=0.12).",
        n_unit="n=21 patients",
        source_pr="#541",
        caveat="IFN family logFC is about 0 in this cohort alone. Four PE captures with n_malignant=0 are why the patient ρ is softer than the sample row (ρ=−0.522).",
        slide_large="Mean log-norm ρ=−0.532 (n=21, p=0.013)",
        slide_lock="%pos ρ=−0.478 (p=0.028)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-3",
        limb="CosMx He2022: CLDN4-high tumor has fewer immune neighbors",
        fit="YES",
        largest_effect="Immune-cell fraction 0.229× at 9 µm (CLDN4 at or above the section 75th percentile of positive counts vs count 0); 8/8 sections and 5/5 patients (#696). Prior count≥1 at 10 µm is 0.377×.",
        locked_primary="Cytotoxic neighbor ratio 0.36 at 50 µm and 0.52 at 100 µm; 8/8 and 5/5; sign P=0.031. Not recomputed.",
        n_unit="8 sections; 5 patients; 765,771 cells in the atlas",
        source_pr="Locked ratios in the public handoff and #698; largest 8/8 fold #696",
        caveat="0.229× does not replace 0.36 / 0.52 (different radius, cut, and immune definition). A raw 0.203× failed the absent-arm floor. Patient-bootstrap winner is median corrected log2 ratio −1.307 (equal-patient ratio of means 0.413) at 10 µm, interval entirely below 0 (#714). At 40 µm, median section means are 1.52 vs 2.69 (median Δ=−1.20), lower in 8/8 (#551). Squidpy median-split ratio 0.883 is 7/8 only (#629).",
        slide_large="0.229× immune fraction at 9 µm; 8/8 and 5/5",
        slide_lock="Cytotoxic 0.36 at 50 µm and 0.52 at 100 µm",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-4",
        limb="CosMx: exclusion, not muzzling, of nearby effectors",
        fit="YES",
        largest_effect="GZMB / PRF1 / NKG7 / IFNG high/low 1.11–1.22; decreased in 0/8 sections.",
        locked_primary="Same. Squidpy nearest-tumor GZMB CPM is higher in 8/8 at 100 µm (median ratio 1.18, #629).",
        n_unit="8 sections; 5 patients",
        source_pr="Public handoff; restated #629 and #738",
        caveat="Write exclusion, not muzzling. No effector gene is lower in 8/8 sections.",
        slide_large="Effector hi/lo 1.11–1.22; 0/8 decreased",
        slide_lock="Same. GZMB nearest-tumor ratio 1.18 (8/8)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-5",
        limb="CosMx: CLDN4-high tumor is less likely to contact a CD8 cell",
        fit="YES",
        largest_effect="FOV-stratified Mantel–Haenszel OR 0.560 (0.547–0.572).",
        locked_primary="Pooled OR 0.584 (0.567–0.602). 5,000 FOV-restricted permutations p=0.00020. OR<1 in 8/8 slides and 5/5 tissues.",
        n_unit="Official 960-plex; overall contact rate 8.9%",
        source_pr="#567",
        caveat="A PanCK-high / CD45-low protein index does not show this deficit. F11R and NECTIN2 are absent from the panel. CDH1 is enriched in the contacts that remain (Δ log-norm +0.053).",
        slide_large="FOV-stratified OR 0.560",
        slide_lock="Pooled OR 0.584; 8/8 and 5/5; perm p=0.00020",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-6",
        limb="CosMx: CLDN4-specific cross-type g(r) vs CD8",
        fit="YES",
        largest_effect="Mean Δg_inhomogeneous = −0.043 at r = 22.5 µm (Wilcoxon p=1.1e-12; Stouffer p=2.0e-10).",
        locked_primary="Same headline. 40/159 FOVs have label-permutation p<0.05; 17 survive BH-FDR.",
        n_unit="159/233 FOVs; 8 samples; 5 patients",
        source_pr="#561",
        caveat="Not detected in Lung5 serial sections. Homogeneous g≪1 versus complete spatial randomness is tumor–stroma geometry and is not a CLDN4-specific claim.",
        slide_large="Δg_inhom = −0.043 at 22.5 µm (p=1.1e-12)",
        slide_lock="Same. 17/159 FOVs survive BH-FDR",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-7",
        limb="Concordant-4 malignant CLDN4-high: IFN and MHC pathways down",
        fit="YES",
        largest_effect="Hallmark IFN-γ NES=−3.87 (patient-bootstrap 95% −4.00 to −2.58). No eligible IFN test in a 310-test grid was more negative (#709).",
        locked_primary="Prior fgsea IFN-γ NES=−3.85 (bootstrap −4.01 to −2.70); IFN-α −3.42; MHC-I −2.75 (#691). Re-fit MHC-I NES=−2.71. Family OLS logFC: chemokine −0.964, IFN −0.584, MHC −0.779 (18 vs 16, #503).",
        n_unit="Expression n=64; Q4 vs Q1 = 18 vs 16",
        source_pr="#709 largest NES; #691 fgsea; #503 family OLS",
        caveat="GSE131907 alone is flat for IFN. Dropping GSE205335, IFN-γ NES falls to −2.40. KEGG tight junction NES=+1.13 and its bootstrap crosses zero — this ranking is not “TJ pathway up.” Largest signature logFC is −0.971; largest canonical gene is CCL5 −2.451.",
        slide_large="IFN-γ NES=−3.87 (grid maximum, #709)",
        slide_lock="Prior fgsea −3.85; IFN-α −3.42; MHC −2.75",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-8",
        limb="CPTAC LSCC: CLDN4 protein vs immune RNA and protein; structural TJ-15 is null",
        fit="YES",
        largest_effect="Eligible composite HLA-A + HLA-C + HLA-F + CD8A protein ρ=−0.531 (n=78, CI −0.680 to −0.333; WES partial −0.471; search permutation p=0.001) (#708).",
        locked_primary="CLDN4 protein vs ImmuneScore ρ=−0.432 (n=78, p=7.9e-5; WES partial −0.298). GEP18 ρ=−0.461. CD8A RNA ρ=−0.437. CD8A protein ρ=−0.444. MHC-I (HLA-A/B/C) ρ=−0.410. TJ-15 protein vs ImmuneScore ρ=−0.082 (n=108, p=0.40).",
        n_unit="CLDN4 quantified in 78/108 LSCC tumors",
        source_pr="#708 largest composite; #289 ImmuneScore and TJ-15; #596 MHC-I and CD8A protein",
        caveat="Treatment-naive surgery, not ICI. Filling the 30 undetected CLDN4 values weakens |ρ| (best filled row −0.474, n=108). The composite is a searched sum; ImmuneScore −0.432 is the pre-specified single-protein row. TJ-15 null on the same freeze is the within-TJ pin.",
        slide_large="Composite protein ρ=−0.531 (n=78, #708)",
        slide_lock="ImmuneScore ρ=−0.432. TJ-15 ρ=−0.082 (null)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-9",
        limb="CLDN4-high malignant signature vs CD8A across bulk LUAD cohorts",
        fit="YES",
        largest_effect="Same winning spec, ImmuneScore meta ρ=−0.574 (I²=0%). That is also the ImmuneScore maximum.",
        locked_primary="CD8A meta ρ=−0.533 (fixed effect, I²=0%, n sum=420, k=4). ssGSEA α=0.75, first 163 genes of the locked 221-gene list, unadjusted. The 221-gene z-mean baseline is ρ=−0.409 (I²=71%).",
        n_unit="OncoSG 169; GSE273377 103 and 60; GSE282774 58; GSE233774 30",
        source_pr="#693",
        caveat="Purity or stromal partial meta ρ=−0.416. The single largest cell (GSE282774 ImmuneScore ρ=−0.690, size 22) is not the cross-study spec. Prefix length was the only size knob; genes were not picked on their CD8 correlation.",
        slide_large="ImmuneScore meta ρ=−0.574 (I²=0%)",
        slide_lock="CD8A meta ρ=−0.533 (n sum=420, I²=0%)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-10",
        limb="Concordant-4: CLDN4-high malignant cells send more barrier ligands to T/NK",
        fit="YES",
        largest_effect="CellPhoneDB expression proportion, outer 10% CLDN4: +27.8 percentage points (mean of F11R, NECTIN2, CDH1, LGALS9); n=60; Wilcoxon p=3.69e-11; 4/4 cohorts positive (#713).",
        locked_primary="Q4 vs Q1 mean +25.6 percentage points (n=64, p=4.91e-12). CellPhoneDB family Δ=+0.253 (expr_prop 0.10, n=53, 52/53 patients positive, #711). LIANA consensus barrier family BH q from 1.3e-9 to 4.9e-8 (#579).",
        n_unit="n=60 patients for the decile; n=64 for Q4 vs Q1",
        source_pr="#713 largest proportion; #711 score Δ; #579 consensus q",
        caveat="Do not quote CellChat probability +0.0037 (#616) as the effect size. The probability-scale searched maximum is +0.191 on n=19. CXCL9/10/11 are not barrier-sized. This is expression, not spatial exclusion.",
        slide_large="+27.8 percentage points (decile; n=60)",
        slide_lock="Q4 vs Q1 +25.6 pp (n=64). CellPhoneDB Δ=+0.253",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-11",
        limb="Concordant-4 Milo: T/NK neighbourhoods down where malignant CLDN4 is high",
        fit="YES",
        largest_effect="Largest median |log2FC| among specs with ≥20 down-hits: 3.102 (24 down, 0 up; k=40, d=30, top 20% vs rest).",
        locked_primary="Primary Milo model is unchanged. Pareto count maximum: 111 T/NK neighbourhoods down and 4 up out of 708 (median |log2FC| 1.596; k=12, d=15, top 40%).",
        n_unit="65 units in every fit",
        source_pr="#701",
        caveat="No specification wins both the count and the |log2FC|. The earlier 65-down row (median |log2FC| 1.991) is inside the grid and off the front.",
        slide_large="Median |log2FC| 3.102 (24 down, 0 up)",
        slide_lock="Front count: 111 down / 4 up (|log2FC| 1.596)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-12",
        limb="OncoSG: CLDN4 vs CD8A / ImmuneScore after PURITY",
        fit="YES",
        largest_effect="Unadjusted ImmuneScore ρ=−0.432 (p=4.59e-9); unadjusted CD8A ρ=−0.416 (p=1.85e-8).",
        locked_primary="PURITY partial ImmuneScore ρ=−0.308 (p=4.90e-5); partial CD8A ρ=−0.285 (p=1.81e-4).",
        n_unit="n=169",
        source_pr="#736",
        caveat="East-Asian surgical LUAD, not ICI.",
        slide_large="Unadj ImmuneScore ρ=−0.432; CD8A −0.416",
        slide_lock="Partial ImmuneScore −0.308; CD8A −0.285",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-13",
        limb="GSE31210: CLDN4 vs CD8A / ImmuneScore",
        fit="PARTIAL",
        largest_effect="Unadjusted ImmuneScore ρ=−0.366 (p=1.50e-8); unadjusted CD8A ρ=−0.341 (p=1.45e-7).",
        locked_primary="ESTIMATE partial ImmuneScore ρ=−0.031 (p=0.639); partial CD8A ρ=−0.127 (p=0.057).",
        n_unit="n=226",
        source_pr="#736",
        caveat="The partial correlations are not significant.",
        slide_large="Unadj ImmuneScore ρ=−0.366; CD8A −0.341",
        slide_lock="ESTIMATE partial −0.031 / −0.127 (ns)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-14",
        limb="TCGA: CLDN4 vs CD8 after keratin adjustment",
        fit="PARTIAL",
        largest_effect="All-patient KRT5/6 sweep: CLDN4 pooled ρ=−0.134 (95% CI −0.194 to −0.073, p=1.68e-5, 8/8 negative) (#705).",
        locked_primary="Pre-specified KRT8/18/19 CD8 mean: CLDN4 pooled ρ=−0.081, beside TACSTD2 −0.069 and CLDN7 −0.109 (#593). Best shared continuous panel (CD3, KRT5/6, no purity): CLDN4 ρ=−0.127 (8/8).",
        n_unit="Eight primary cohorts",
        source_pr="#593 primary; #705 all-patient maximum",
        caveat="Effects are small, and the sweep inflates them. The surface-gene screen does not pin CLDN4: LUAD rank 7, partial +0.122, with CLDN7 / EPCAM / MUC1 stronger (public handoff). BRCA keratin caveat on the TACSTD2 screen is −0.071.",
        slide_large="KRT5/6 sweep ρ=−0.134 (8/8)",
        slide_lock="KRT8/18/19 ρ=−0.081. Surface rank does not pin CLDN4",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-15",
        limb="Concordant-4: CLDN4–T/NK association remains after TACSTD2",
        fit="YES",
        largest_effect="CLDN4 %pos vs T/NK, partial on TACSTD2: ρ=−0.543 (p=2.0e-5, I²=0%, negative in 4/4).",
        locked_primary="Unadjusted ρ=−0.531. Two-stage residual after TACSTD2 is r=−0.460 (p=1.2e-4, #725). SEM direct path −0.525 (95% CI −0.726 to −0.264).",
        n_unit="N=65",
        source_pr="#733 and #718; residual #725",
        caveat="TACSTD2 %pos vs T/NK is itself null (ρ=−0.112, p=0.46). The immune association sits on CLDN4. This is not a claim that TACSTD2 cold is mediated by CLDN4: the concordant-4 and CosMx mediation matrix is null (#718), and no grid row kept a negative direct TACSTD2 coefficient with a mediation proportion inside (0, 1) (#733).",
        slide_large="Partial ρ=−0.543 given TACSTD2 (4/4, I²=0)",
        slide_lock="Unadjusted ρ=−0.531. TACSTD2 alone is −0.112 (null)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-16",
        limb="Public scRNA: CLDN4 vs CLDN7 / EPCAM on the same T/NK endpoint",
        fit="PARTIAL",
        largest_effect="After keratin and CLDN7 are both partialled, CLDN4 unique ρ=−0.465 (permutation p=0.0002). CLDN7 given CLDN4 is ρ=−0.184 (p=0.15). EPCAM's unique partial collapses (ρ=−0.013).",
        locked_primary="Keratin-partial %pos vs T/NK: CLDN4 −0.478 (p=2.7e-4, I²=0), EPCAM −0.368, CLDN7 −0.352 (p=0.10, I²=56%), CLDN1 +0.225. Head-to-head CLDN4 vs CLDN7 Δ=−0.206, permutation p=0.17; vs EPCAM p=0.18.",
        n_unit="N=65 concordant-4 units",
        source_pr="#644",
        caveat="Public data do not separate CLDN4 from CLDN7 by a difference test. Inside GSE131907 the CLDN7 keratin partial is stronger (−0.658 vs −0.514). Dropping GSE205335 reverses the meta rank. TCGA keratin CD8 ρ moves from −0.069 to −0.040 after CLDN4 (share 41.3%), but the CLDN4-minus-CLDN7 share interval contains 0, and CLDN4's further share after CLDN7 and EPCAM is 1.8% (#728). Where TACSTD2–T/NK is negative, CLDN7 attenuates it about as much as CLDN4 (#732). The stable pin is the private KD co-culture, which is not in this repository.",
        slide_large="Unique ρ=−0.465 after keratin and CLDN7 (p=0.0002)",
        slide_lock="Head-to-head vs CLDN7 p=0.17 (not separated)",
    ),
    dict(
        part="Part 2 — CLDN4 pin",
        limb_id="P2-17",
        limb="SEM: graph TACSTD2 → CLDN4 → T/NK on concordant-4",
        fit="PARTIAL",
        largest_effect="Concordant-4 BIC prefers TACSTD2 → CLDN4 → immune by ΔBIC 14.66 versus the reverse order and 17.95 versus independence. Indirect path −0.263 (95% CI −0.439 to −0.123).",
        locked_primary="Direct CLDN4 path −0.525 (CI −0.726 to −0.264). Total TACSTD2 effect −0.151 (CI −0.394 to +0.136) includes zero. TCGA seven-cohort BIC sum also prefers this graph (ΔBIC 11.60) with pooled indirect −0.022 and a mediation ratio 42.1% whose CI is 19.5 to 114.9.",
        n_unit="Concordant-4 n=65; TCGA seven-cohort sum n=3,444",
        source_pr="#725",
        caveat="Do not quote the 174% path ratio as a mediation share. GSE205335 alone prefers independence. TCGA cohorts do not vote as one graph, and adding LUSC flips the sum. A lower BIC is a covariance description, not a knockdown.",
        slide_large="ΔBIC 14.66 for TACSTD2 → CLDN4 → T/NK",
        slide_lock="CLDN4 direct −0.525. TACSTD2 total CI includes 0",
    ),
]

SKIPS = [
    ("Concordant-4 TACSTD2 %pos vs T/NK", "NULL", "ρ=−0.112, p=0.46, I²=15%, N=65. Do not quote CLDN4 ρ=−0.531 as TACSTD2.", "#718 / #739"),
    ("CosMx TACSTD2 at 50 and 100 µm", "Not 8/8", "Count ratios 0.854 (7/8) and 0.947 (4/8). Short-range 0.455 / 0.670 is the YES row.", "#726"),
    ("CPTAC TROP2 protein vs CD8A protein", "NO", "LUAD ρ=−0.038 and LSCC ρ=−0.103; both CIs cross 0.", "#721"),
    ("CPTAC LSCC TROP2 protein vs xCell CD8", "NULL", "ρ=−0.080, n=108, p=0.41.", "#99"),
    ("TISMO Tacstd2 after ICB", "Not fewer T/NK", "49/64 groups up, Wilcoxon p=5.8e-5. Cldn4 is 34/64, not 49/64. Baseline CD8 link is positive. No true KL/KP lung line.", "#152 / #173"),
    ("Public GEMM Cldn4 %pos vs T/NK", "Opposite", "REML pooled ρ=+0.477, p=0.302, I²=72%, 34 mice. IFN pooled ρ=+0.092.", "#677"),
    ("Mouse n=4 Cldn4 or Tacstd2 vs T fraction", "Exploratory", "ρ=−1.000, exact p=0.083. Not confirmatory. IFN/APM on the same grid can be ρ=+1.", "#700 / #743"),
    ("DepMap NSCLC TROP2 vs Hallmark IFN-γ", "Opposite of cold", "ρ=+0.339 (n=143). TJ does not explain it away (partial +0.277).", "#740"),
    ("Concordant-4 broad TJ score vs T/NK", "NULL", "ρ=+0.137, p=0.60. TJ-core ρ=−0.191, p=0.41. Immune exclusion stays on CLDN4 %pos.", "#741"),
    ("CLDN4 Q4 vs Q1 KEGG tight junction", "Not up", "NES=+1.13; patient bootstrap crosses zero.", "#691"),
    ("TACSTD2 → CLDN4 → immune mediation", "NULL as a proportion", "CosMx and concordant-4 primary mediation rows are null (#718). No spec kept mediation proportion inside (0, 1) with a negative direct coefficient (#733). CosMx 10–20 µm TACSTD2 exclusion is not accounted for by CLDN4 (rank attenuation 0.124 / 0.103, #726).", "#718 / #733 / #726"),
    ("TCGA surface ranking as a CLDN4 pin", "Does not pin", "LUAD CLDN4 rank 7, partial +0.122. CLDN7, EPCAM, and MUC1 are stronger. Head-to-head vs CLDN7 on concordant-4 is p=0.17 (#644).", "Handoff; #644; #728"),
    ("GSE126044 non-responders have higher TJ", "Fragile", "Recovered only for one 7-gene z-mean (p=0.019, n=5 vs 11). CLDN4 alone p=0.115. Response is aliased with FFPE.", "#90"),
    ("Visium / Stereo-seq / GeoMx / official Xenium lung panel", "Do not upgrade", "Mixed spots, missing CLDN4, or opposite sign. Do not write a same-spot correlation as spatial exclusion.", "Handoff; #544–#566 except #551 and #561"),
    ("Human ICI bulk CLDN4 rises after resistance", "Not supported", "GSE126044, GSE135222, GSE248249, GSE248378 and related sets are small, mixed, or opposite.", "Handoff"),
    ("Public mouse merged with private 8KL", "Do not merge", "GSE165641, GSE180963, GSE154977 stay separate from the private eight KL matrices.", "Handoff"),
    ("Discordant human scRNA added to concordant-4", "Do not add", "GSE148071, GSE127465, GSE154826, GSE200563, E-MTAB-13526 dilute or flip the sign.", "Handoff"),
    ("Broad Tacstd2-high GSEA, TJ in ranks 1–3", "NO", "any_cohort_highlight_in_top3 = false on a 589-set universe (#744). The focused-set rank 1 (#746) stays PARTIAL in P1-T7.", "#744"),
]


def write_tsv() -> None:
    fields = [
        "part",
        "limb_id",
        "limb",
        "fit",
        "largest_effect",
        "locked_primary",
        "n_unit",
        "source_pr",
        "caveat",
    ]
    path = OUT / "master_table.tsv"
    with path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore", delimiter="\t")
        w.writeheader()
        for row in ROWS:
            w.writerow(row)
    skip_path = OUT / "skips.tsv"
    with skip_path.open("w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["item", "call", "number", "source"])
        w.writerows(SKIPS)


def md_cell(text: str) -> str:
    return text.replace("|", "\\|").replace("\n", " ")


def write_md() -> None:
    n_yes = sum(r["fit"] == "YES" for r in ROWS)
    n_partial = sum(r["fit"] == "PARTIAL" for r in ROWS)
    lines = []
    lines.append("# Master PPT table — story-fit locked results")
    lines.append("")
    lines.append("Public PRs only. No new fits. No private 8KL matrices. Every coefficient is copied from the cited PR or from the public handoff.")
    lines.append("")
    lines.append(f"**{len(ROWS)} limbs: {n_yes} YES, {n_partial} PARTIAL.** Part 1 is Tacstd2 / TROP2 immune-cold plus Tacstd2 tracking tight junction. Part 2 is the CLDN4 pin. NO, NULL, and opposite results are in the skip list and are not slide support.")
    lines.append("")
    lines.append("**Largest effect** is the largest eligible story-direction number the source PR reported for that limb. **Locked primary** is the pre-specified number when a later sweep is larger. A searched maximum is not a second confirmatory test. Searched p-values are descriptive unless that PR reported a permutation that included the search.")
    lines.append("")
    lines.append("Paste `MASTER_STORY_FIT.pptx` if you need slides. This markdown file is the full table, including caveats that do not fit on a slide.")
    lines.append("")
    for part in dict.fromkeys(r["part"] for r in ROWS):
        lines.append(f"## {part}")
        lines.append("")
        lines.append("| ID | Limb | Fit | Largest effect | Locked primary | n | PR |")
        lines.append("|---|---|---|---|---|---|---|")
        for r in ROWS:
            if r["part"] != part:
                continue
            lines.append(
                "| {limb_id} | {limb} | {fit} | {largest_effect} | {locked_primary} | {n_unit} | {source_pr} |".format(
                    **{k: md_cell(r[k]) for k in ("limb_id", "limb", "fit", "largest_effect", "locked_primary", "n_unit", "source_pr")}
                )
            )
        lines.append("")
        lines.append("| ID | Caveat (do not drop) |")
        lines.append("|---|---|")
        for r in ROWS:
            if r["part"] != part:
                continue
            lines.append(f"| {r['limb_id']} | {md_cell(r['caveat'])} |")
        lines.append("")
    lines.append("## Do not paste as support")
    lines.append("")
    lines.append("| Item | Call | Number | Source |")
    lines.append("|---|---|---|---|")
    for item, call, number, source in SKIPS:
        lines.append(f"| {md_cell(item)} | {md_cell(call)} | {md_cell(number)} | {md_cell(source)} |")
    lines.append("")
    lines.append("## What this table does not do")
    lines.append("")
    lines.append("- It does not recompute CosMx ratios 0.36 / 0.52 or concordant-4 ρ=−0.531.")
    lines.append("- It does not treat a same-spot Visium correlation as spatial exclusion.")
    lines.append("- It does not merge public mouse Harmony objects with the private eight KL matrices.")
    lines.append("- It does not claim TISMO 49/64 is immune exclusion. That lock is Tacstd2 up after ICB.")
    lines.append("- It does not claim public data uniquely rank CLDN4 over CLDN7. P2-16 is PARTIAL on purpose.")
    lines.append("")
    (OUT / "MASTER_PPT_TABLE.md").write_text("\n".join(lines))


def set_run(paragraph, text, size, bold=False, color=None, font="Calibri"):
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color or RGBColor(0x1F, 0x29, 0x37)
    return run


def shade_cell_xml(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("a:solidFill"):
            tcPr.remove(child)
    solid = tcPr.makeelement(qn("a:solidFill"), {})
    srgb = solid.makeelement(qn("a:srgbClr"), {"val": hex_color})
    solid.append(srgb)
    tcPr.append(solid)


def write_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Master PPT table — story-fit YES and PARTIAL limbs"
    prs.core_properties.subject = "Tacstd2 cold + TJ; CLDN4 pin. Numbers from existing PRs only."

    navy = RGBColor(0x0F, 0x2C, 0x4C)
    ink = RGBColor(0x1F, 0x29, 0x37)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    mute = RGBColor(0x4B, 0x55, 0x63)

    def blank():
        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0xF7, 0xF5, 0xF0)
        return slide

    def add_title(slide, text, sub):
        box = slide.shapes.add_textbox(Inches(0.35), Inches(0.16), Inches(12.6), Inches(0.38))
        p = box.text_frame.paragraphs[0]
        set_run(p, text, 18, bold=True, color=navy)
        box2 = slide.shapes.add_textbox(Inches(0.35), Inches(0.50), Inches(12.6), Inches(0.28))
        p2 = box2.text_frame.paragraphs[0]
        set_run(p2, sub, 11, color=mute)

    # Title
    slide = blank()
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.12))  # mso shape rectangle = 1
    bar.fill.solid()
    bar.fill.fore_color.rgb = navy
    bar.line.fill.background()
    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.45), Inches(12.3), Inches(1.1))
    tf = title.text_frame
    tf.word_wrap = True
    set_run(tf.paragraphs[0], "Master table — story-fit results", 32, bold=True, color=navy)
    p = tf.add_paragraph()
    set_run(p, "Part 1  Tacstd2 cold  +  tight junction          Part 2  CLDN4 pin", 18, color=ink)
    n_yes = sum(r["fit"] == "YES" for r in ROWS)
    n_partial = sum(r["fit"] == "PARTIAL" for r in ROWS)
    body = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(12.2), Inches(4.6))
    btf = body.text_frame
    btf.word_wrap = True
    bullets = [
        f"{len(ROWS)} limbs on the following slides: {n_yes} YES and {n_partial} PARTIAL. Copied from existing PRs. No new fits.",
        "Largest effect = the biggest eligible number in the story direction. Locked primary = the pre-specified number when a sweep is larger. The lock is not replaced.",
        "Part 1 cold is short-range CosMx, OncoSG, and TCGA CD8 — not concordant-4 TACSTD2 (that ρ is −0.112).",
        "Part 1 TJ is real on GSE137244 (Cldn4 Δ=+5.570), concordant-4 apical-junction rank 1, and TCGA/GSE31210 GSEA. A 589-set ranking does not put TJ in the top 3.",
        "Part 2 pin: concordant-4 ρ=−0.531 (largest scVI |ρ| −0.736), CosMx 0.36 / 0.52 (largest 8/8 fold 0.229× at 9 µm), IFN-γ NES −3.87, LSCC protein ρ=−0.531.",
        "Public data do not separate CLDN4 from CLDN7 (difference p=0.17). That row stays PARTIAL. Private KD co-culture is outside this repo.",
        "Skip list is the last slide. TISMO 49/64 is Tacstd2 up after ICB, not fewer T cells.",
    ]
    for i, text in enumerate(bullets):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        set_run(p, text, 16, color=ink)

    headers = ["ID", "Limb", "Fit", "Largest effect", "Locked primary", "PR"]
    widths = [0.72, 2.55, 0.95, 4.15, 3.15, 1.35]
    groups = []
    current = None
    bucket = []
    for row in ROWS:
        if row["part"] != current:
            if bucket:
                groups.append((current, bucket))
            current = row["part"]
            bucket = [row]
        else:
            bucket.append(row)
    if bucket:
        groups.append((current, bucket))

    # Split long parts into chunks of 7
    slides_spec = []
    for part, rows in groups:
        for i in range(0, len(rows), 7):
            chunk = rows[i : i + 7]
            page = i // 7 + 1
            pages = (len(rows) + 6) // 7
            label = part if pages == 1 else f"{part}  ({page}/{pages})"
            slides_spec.append((label, chunk))

    fit_color = {"YES": "1F7A4D", "PARTIAL": "8A5A00"}
    fit_bg = {"YES": "E5F4EC", "PARTIAL": "FFF4D6"}

    for label, chunk in slides_spec:
        slide = blank()
        add_title(
            slide,
            label,
            "Largest effect is not a replacement for the locked primary. Caveats are in MASTER_PPT_TABLE.md.",
        )
        table_shape = slide.shapes.add_table(
            len(chunk) + 1,
            len(headers),
            Inches(0.28),
            Inches(0.88),
            Inches(12.78),
            Inches(6.35),
        )
        table = table_shape.table
        for j, w in enumerate(widths):
            table.columns[j].width = Inches(w)
        for j, h in enumerate(headers):
            cell = table.cell(0, j)
            shade_cell_xml(cell, "0F2C4C")
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            set_run(p, h, 10, bold=True, color=white)
        for i, row in enumerate(chunk, start=1):
            vals = [
                row["limb_id"],
                row["limb"],
                row["fit"],
                row["slide_large"],
                row["slide_lock"],
                row["source_pr"],
            ]
            bg = "FFFFFF" if i % 2 else "F3F1EB"
            for j, val in enumerate(vals):
                cell = table.cell(i, j)
                shade_cell_xml(cell, fit_bg[row["fit"]] if j == 2 else bg)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                color = RGBColor.from_string(fit_color[row["fit"]]) if j == 2 else ink
                set_run(p, val, 9, bold=(j in (0, 2)), color=color)
        notes = slide.notes_slide.notes_text_frame
        notes.text = "\n".join(f"{r['limb_id']}: {r['caveat']}" for r in chunk)

    # Skips slide — two columns of short lines would overflow. Use a table of the
    # items that most often get upgraded by mistake.
    slide = blank()
    add_title(slide, "Do not paste as support", "These calls are real and are not slide evidence for either part.")
    show = SKIPS[:12]
    headers2 = ["Item", "Call", "Number", "Source"]
    widths2 = [3.3, 1.5, 6.3, 1.7]
    table_shape = slide.shapes.add_table(
        len(show) + 1, 4, Inches(0.28), Inches(0.88), Inches(12.78), Inches(6.35)
    )
    table = table_shape.table
    for j, w in enumerate(widths2):
        table.columns[j].width = Inches(w)
    for j, h in enumerate(headers2):
        cell = table.cell(0, j)
        shade_cell_xml(cell, "6B2D2D")
        tf = cell.text_frame
        tf.word_wrap = True
        set_run(tf.paragraphs[0], h, 10, bold=True, color=white)
    for i, (item, call, number, source) in enumerate(show, start=1):
        # shorten number for the slide
        short = number if len(number) < 180 else number[:177] + "…"
        vals = [item, call, short, source]
        for j, val in enumerate(vals):
            cell = table.cell(i, j)
            shade_cell_xml(cell, "F8EAEA" if i % 2 else "FFFFFF")
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            set_run(tf.paragraphs[0], val, 8, bold=(j == 1), color=ink)
    notes = slide.notes_slide.notes_text_frame
    notes.text = "Full skip list is in MASTER_PPT_TABLE.md. Remaining items: " + "; ".join(
        s[0] for s in SKIPS[12:]
    )

    prs.save(OUT / "MASTER_STORY_FIT.pptx")


def write_provenance() -> None:
    payload = {
        "rule": "Numbers are copied from existing PR bodies or the public handoff. No new fits.",
        "n_limbs": len(ROWS),
        "n_yes": sum(r["fit"] == "YES" for r in ROWS),
        "n_partial": sum(r["fit"] == "PARTIAL" for r in ROWS),
        "parts": {
            "part1_cold": "Tacstd2 / TROP2 immune-cold. Concordant-4 TACSTD2 composition is NULL and is not included.",
            "part1_tj": "Tacstd2-high tracks CLDN4 / tight junction / keratin. Broad GSEA does not rank TJ #1–3.",
            "part2_cldn4_pin": "CLDN4, not TACSTD2 and not structural TJ-15 protein, carries the immune-cold association. Uniqueness versus CLDN7 is PARTIAL.",
        },
        "limbs": [
            {"id": r["limb_id"], "fit": r["fit"], "pr": r["source_pr"]} for r in ROWS
        ],
    }
    (OUT / "provenance.json").write_text(json.dumps(payload, indent=2) + "\n")


def main() -> None:
    write_tsv()
    write_md()
    write_provenance()
    write_pptx()
    print(f"wrote {len(ROWS)} limbs to {OUT}")


if __name__ == "__main__":
    main()
