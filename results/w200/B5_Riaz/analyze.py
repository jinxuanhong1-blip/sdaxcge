#!/usr/bin/env python3
"""Open melanoma ICI RNA: CLDN4 versus response, including Riaz 2017."""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import math
import re
import urllib.request
import zipfile
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from scipy import stats


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
CLDN4_ENTREZ = "1364"
SEED = 91061

URLS = {
    "GSE91061_fpkm.csv.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE91nnn/GSE91061/suppl/"
        "GSE91061_BMS038109Sample.hg19KnownGene.fpkm.csv.gz"
    ),
    "GSE91061_series_matrix.txt.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE91nnn/GSE91061/matrix/"
        "GSE91061_series_matrix.txt.gz"
    ),
    "GSE78220_PatientFPKM.xlsx": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE78nnn/GSE78220/suppl/"
        "GSE78220_PatientFPKM.xlsx"
    ),
    "GSE78220_series_matrix.txt.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE78nnn/GSE78220/matrix/"
        "GSE78220_series_matrix.txt.gz"
    ),
    "GSE115821-GPL11154_series_matrix.txt.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE115nnn/GSE115821/matrix/"
        "GSE115821-GPL11154_series_matrix.txt.gz"
    ),
    "GSE115821-GPL18573_series_matrix.txt.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE115nnn/GSE115821/matrix/"
        "GSE115821-GPL18573_series_matrix.txt.gz"
    ),
    "GSE115821_MGH_counts.csv.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE115nnn/GSE115821/suppl/"
        "GSE115821_MGH_counts.csv.gz"
    ),
    "GSE93157_series_matrix.txt.gz": (
        "https://ftp.ncbi.nlm.nih.gov/geo/series/GSE93nnn/GSE93157/matrix/"
        "GSE93157_series_matrix.txt.gz"
    ),
    "liu_MOESM4.xlsx": (
        "https://static-content.springer.com/esm/"
        "art%3A10.1038%2Fs41591-019-0654-5/MediaObjects/"
        "41591_2019_654_MOESM4_ESM.xlsx"
    ),
    "liu_addData.zip": (
        "https://raw.githubusercontent.com/vanallenlab/schadendorf-pd1/"
        "master/data/addData.zip"
    ),
}


def download_inputs() -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    for name, url in URLS.items():
        path = DATA / name
        if not path.exists():
            print(f"Downloading {url}")
            urllib.request.urlretrieve(url, path)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_geo_sample_rows(path: Path) -> list[list[str]]:
    rows: list[list[str]] = []
    with gzip.open(path, "rt") as handle:
        for line in handle:
            if line.startswith("!Sample_"):
                rows.append(next(csv.reader([line], delimiter="\t")))
    return rows


def geo_one(rows: list[list[str]], name: str) -> list[str]:
    matches = [row[1:] for row in rows if row[0] == name]
    if len(matches) != 1:
        raise ValueError(f"{name}: expected 1 row, found {len(matches)}")
    return matches[0]


def geo_characteristics(rows: list[list[str]]) -> dict[str, list[str]]:
    mapping: dict[str, list[str]] = {}
    for values in (row[1:] for row in rows if row[0] == "!Sample_characteristics_ch1"):
        key = values[0].split(":", 1)[0]
        mapping[key] = [value.split(": ", 1)[1] if ": " in value else value for value in values]
    return mapping


def write_csv(path: Path, rows: list[dict], fieldnames: list[str] | None = None) -> None:
    if not rows:
        raise ValueError(f"No rows for {path}")
    names = fieldnames or list(rows[0])
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=names, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def bootstrap_auc_ci(responders: np.ndarray, nonresponders: np.ndarray, draws: int = 20_000) -> tuple[float, float]:
    rng = np.random.default_rng(SEED)
    estimates = np.empty(draws)
    for i in range(draws):
        a = rng.choice(responders, len(responders), replace=True)
        b = rng.choice(nonresponders, len(nonresponders), replace=True)
        estimates[i] = stats.mannwhitneyu(a, b).statistic / (len(a) * len(b))
    return tuple(np.quantile(estimates, [0.025, 0.975]))


def median_split_or(values: np.ndarray, responder: np.ndarray) -> dict[str, float | int]:
    median = float(np.median(values))
    high = values >= median
    a = int(np.sum(high & responder))
    b = int(np.sum(high & ~responder))
    c = int(np.sum(~high & responder))
    d = int(np.sum(~high & ~responder))
    fisher = stats.fisher_exact([[a, b], [c, d]], alternative="two-sided")
    or_ha = ((a + 0.5) * (d + 0.5)) / ((b + 0.5) * (c + 0.5))
    se = math.sqrt(1 / (a + 0.5) + 1 / (b + 0.5) + 1 / (c + 0.5) + 1 / (d + 0.5))
    log_or = math.log(or_ha)
    return {
        "median_split": median,
        "high_responder": a,
        "high_nonresponder": b,
        "low_responder": c,
        "low_nonresponder": d,
        "or_cldn4high_response": or_ha,
        "or_log": log_or,
        "or_se_log": se,
        "or_ci_low": math.exp(log_or - 1.96 * se),
        "or_ci_high": math.exp(log_or + 1.96 * se),
        "or_fisher_p": float(fisher.pvalue),
    }


def summarize_expression(values: np.ndarray, responder: np.ndarray) -> dict[str, float | int]:
    r = values[responder]
    nr = values[~responder]
    test = stats.mannwhitneyu(r, nr, alternative="two-sided", method="asymptotic")
    auc = float(test.statistic / (len(r) * len(nr)))
    auc_lo, auc_hi = bootstrap_auc_ci(r, nr)
    out = {
        "n_responder": int(len(r)),
        "n_nonresponder": int(len(nr)),
        "median_responder": float(np.median(r)),
        "median_nonresponder": float(np.median(nr)),
        "mann_whitney_u": float(test.statistic),
        "mann_whitney_p": float(test.pvalue),
        "auc_responder_higher": auc,
        "auc_bootstrap_95ci_low": auc_lo,
        "auc_bootstrap_95ci_high": auc_hi,
    }
    out.update(median_split_or(values, responder))
    return out


def mantel_haenszel(rows: list[dict]) -> dict[str, float]:
    num = 0.0
    den = 0.0
    var_num = 0.0
    for row in rows:
        a = row["high_responder"] + 0.5
        b = row["high_nonresponder"] + 0.5
        c = row["low_responder"] + 0.5
        d = row["low_nonresponder"] + 0.5
        n = a + b + c + d
        num += a * d / n
        den += b * c / n
        var_num += (
            ((a + b) * (c + d) * (a + c) * (b + d)) / ((n**2) * (n - 1))
        )
    or_mh = num / den
    se = math.sqrt(var_num) / den
    log_or = math.log(or_mh)
    return {
        "or_mh": or_mh,
        "or_ci_low": math.exp(log_or - 1.96 * se),
        "or_ci_high": math.exp(log_or + 1.96 * se),
    }


def load_riaz() -> tuple[list[dict], dict[str, dict], dict[str, np.ndarray]]:
    rows = parse_geo_sample_rows(DATA / "GSE91061_series_matrix.txt.gz")
    titles = geo_one(rows, "!Sample_title")
    accessions = geo_one(rows, "!Sample_geo_accession")
    chars = geo_characteristics(rows)
    if len(titles) != 109:
        raise ValueError("Unexpected Riaz sample count")
    metadata = {}
    for title, gsm, visit, response in zip(
        titles, accessions, chars["visit (pre or on treatment)"], chars["response"], strict=True
    ):
        match = re.match(r"Pt([^_]+)_", title)
        if not match:
            raise ValueError(title)
        metadata[title] = {
            "sample": title,
            "gsm": gsm,
            "patient": f"Pt{match.group(1)}",
            "visit": visit,
            "response_raw": response,
        }
    with gzip.open(DATA / "GSE91061_fpkm.csv.gz", "rt", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        matches = [row for row in reader if row and row[0] == CLDN4_ENTREZ]
    if len(matches) != 1:
        raise ValueError("Riaz CLDN4 row missing")
    expression = dict(zip(header[1:], (float(v) for v in matches[0][1:]), strict=True))
    sample_rows = []
    for sample, fpkm in expression.items():
        meta = metadata[sample]
        recist = meta["response_raw"]
        group = "Responder" if recist == "PRCR" else "Nonresponder" if recist in {"SD", "PD"} else ""
        sample_rows.append(
            {
                **meta,
                "cohort": "Riaz2017_GSE91061",
                "cldn4_value": fpkm,
                "cldn4_log2": math.log2(fpkm + 1),
                "binary_group": group,
            }
        )
    sample_rows.sort(key=lambda row: (row["patient"], row["visit"]))
    eligible = [row for row in sample_rows if row["visit"] == "Pre" and row["binary_group"]]
    values = np.array([row["cldn4_log2"] for row in eligible])
    responder = np.array([row["binary_group"] == "Responder" for row in eligible])
    summary = summarize_expression(values, responder)
    groups = {
        "Responder": values[responder],
        "Nonresponder": values[~responder],
    }
    return sample_rows, {"pretreatment": summary}, groups


def riaz_exploratory(sample_rows: list[dict]) -> tuple[list[dict], dict[str, dict], dict[str, dict[str, np.ndarray]]]:
    eligible = [row for row in sample_rows if row["binary_group"]]
    on_rows = [row for row in eligible if row["visit"] == "On"]
    on_values = np.array([row["cldn4_log2"] for row in on_rows])
    on_resp = np.array([row["binary_group"] == "Responder" for row in on_rows])
    by_patient: dict[str, dict[str, dict]] = {}
    for row in eligible:
        by_patient.setdefault(row["patient"], {})[row["visit"]] = row
    change_rows = []
    for patient, visits in by_patient.items():
        if "Pre" not in visits or "On" not in visits:
            continue
        pre, on = visits["Pre"], visits["On"]
        if pre["response_raw"] != on["response_raw"]:
            raise ValueError(f"Inconsistent Riaz response for {patient}")
        change_rows.append(
            {
                "patient": patient,
                "response": pre["response_raw"],
                "binary_group": pre["binary_group"],
                "pre_log2": pre["cldn4_log2"],
                "on_log2": on["cldn4_log2"],
                "change_on_minus_pre": on["cldn4_log2"] - pre["cldn4_log2"],
            }
        )
    change_values = np.array([row["change_on_minus_pre"] for row in change_rows])
    change_resp = np.array([row["binary_group"] == "Responder" for row in change_rows])
    summaries = {
        "on_treatment": summarize_expression(on_values, on_resp),
        "paired_change": summarize_expression(change_values, change_resp),
    }
    groups = {
        "on_treatment": {
            "Responder": on_values[on_resp],
            "Nonresponder": on_values[~on_resp],
        },
        "paired_change": {
            "Responder": change_values[change_resp],
            "Nonresponder": change_values[~change_resp],
        },
    }
    return change_rows, summaries, groups


def load_hugo() -> tuple[list[dict], dict]:
    rows = parse_geo_sample_rows(DATA / "GSE78220_series_matrix.txt.gz")
    titles = geo_one(rows, "!Sample_title")
    accessions = geo_one(rows, "!Sample_geo_accession")
    chars = geo_characteristics(rows)
    response = chars["anti-pd-1 response"]
    wb = load_workbook(DATA / "GSE78220_PatientFPKM.xlsx", read_only=True, data_only=True)
    ws = wb["FPKM"]
    iterator = ws.iter_rows(values_only=True)
    header = next(iterator)
    cldn4 = None
    for row in iterator:
        if row and row[0] == "CLDN4":
            cldn4 = row
            break
    wb.close()
    if cldn4 is None:
        raise ValueError("Hugo CLDN4 missing")
    expr = {str(name).split(".", 1)[0]: float(value) for name, value in zip(header[1:], cldn4[1:], strict=True)}
    sample_rows = []
    for title, gsm, recist in zip(titles, accessions, response, strict=True):
        group = (
            "Responder"
            if recist in {"Complete Response", "Partial Response"}
            else "Nonresponder"
            if recist == "Progressive Disease"
            else ""
        )
        fpkm = expr[title]
        sample_rows.append(
            {
                "cohort": "Hugo2016_GSE78220",
                "sample": title,
                "gsm": gsm,
                "patient": title,
                "visit": "Pre",
                "response_raw": recist,
                "cldn4_value": fpkm,
                "cldn4_log2": math.log2(fpkm + 1),
                "binary_group": group,
            }
        )
    eligible = [row for row in sample_rows if row["binary_group"]]
    values = np.array([row["cldn4_log2"] for row in eligible])
    responder = np.array([row["binary_group"] == "Responder" for row in eligible])
    return sample_rows, summarize_expression(values, responder)


def _parse_auslander_matrix(path: Path) -> list[dict]:
    rows = parse_geo_sample_rows(path)
    titles = geo_one(rows, "!Sample_title")
    accessions = geo_one(rows, "!Sample_geo_accession")
    chars = geo_characteristics(rows)
    out = []
    for i, title in enumerate(titles):
        out.append(
            {
                "sample": title,
                "gsm": accessions[i],
                "visit_raw": chars["treatment state"][i],
                "antibody": chars["antibody"][i],
                "response_raw": chars["response"][i],
                "patient": chars["patient id"][i],
            }
        )
    return out


def load_auslander() -> tuple[list[dict], dict]:
    meta = _parse_auslander_matrix(DATA / "GSE115821-GPL11154_series_matrix.txt.gz")
    meta += _parse_auslander_matrix(DATA / "GSE115821-GPL18573_series_matrix.txt.gz")
    with gzip.open(DATA / "GSE115821_MGH_counts.csv.gz", "rt", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        header[0] = header[0].lstrip("\ufeff")
        cldn4 = next(row for row in reader if row and row[0] == "CLDN4")
    counts = {name: float(value) for name, value in zip(header[6:], cldn4[6:], strict=True)}
    libsize = {name: 0.0 for name in counts}
    with gzip.open(DATA / "GSE115821_MGH_counts.csv.gz", "rt", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        for row in reader:
            for name, value in zip(header[6:], row[6:], strict=True):
                libsize[name] += float(value)

    def column_for(title: str) -> str:
        if title in counts:
            return title
        bam = f"{title}.bam"
        if bam in counts:
            return bam
        raise KeyError(title)

    sample_rows = []
    for item in meta:
        col = column_for(item["sample"])
        cpm = 1e6 * counts[col] / libsize[col] if libsize[col] else 0.0
        dabrafenib = "dabrafenib" in item["visit_raw"].lower()
        pre = item["visit_raw"].startswith("PRE") and not dabrafenib
        group = (
            "Responder"
            if item["response_raw"] == "R"
            else "Nonresponder"
            if item["response_raw"] == "NR"
            else ""
        )
        sample_rows.append(
            {
                "cohort": "Auslander2018_GSE115821",
                "sample": item["sample"],
                "gsm": item["gsm"],
                "patient": item["patient"],
                "visit": "Pre" if pre else item["visit_raw"],
                "antibody": item["antibody"],
                "response_raw": item["response_raw"],
                "cldn4_value": cpm,
                "cldn4_log2": math.log2(cpm + 1),
                "binary_group": group,
                "count_column": col,
            }
        )
    by_patient: dict[str, list[dict]] = {}
    for row in sample_rows:
        if row["visit"] == "Pre" and row["binary_group"]:
            by_patient.setdefault(row["patient"], []).append(row)
    collapsed = []
    for patient, items in sorted(by_patient.items()):
        logs = [item["cldn4_log2"] for item in items]
        collapsed.append(
            {
                "cohort": "Auslander2018_GSE115821",
                "sample": ";".join(item["sample"] for item in items),
                "gsm": ";".join(item["gsm"] for item in items),
                "patient": patient,
                "visit": "Pre",
                "antibody": ";".join(sorted({item["antibody"] for item in items})),
                "response_raw": items[0]["response_raw"],
                "cldn4_value": float(np.mean([item["cldn4_value"] for item in items])),
                "cldn4_log2": float(np.mean(logs)),
                "binary_group": items[0]["binary_group"],
                "n_pre_samples": len(items),
            }
        )
    values = np.array([row["cldn4_log2"] for row in collapsed])
    responder = np.array([row["binary_group"] == "Responder" for row in collapsed])
    return sample_rows, collapsed, summarize_expression(values, responder)


def load_liu() -> tuple[list[dict], dict]:
    wb = load_workbook(DATA / "liu_MOESM4.xlsx", read_only=True, data_only=True)
    ws = wb["Supplemental Table 1"]
    table_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = table_rows[2]
    clinical = {}
    for row in table_rows[3:]:
        if not row or not row[0] or not str(row[0]).startswith("Patient"):
            continue
        record = {str(name): value for name, value in zip(header, row, strict=True) if name}
        record["patient"] = str(row[0])
        clinical[record["patient"]] = record
    with zipfile.ZipFile(DATA / "liu_addData.zip") as archive:
        with archive.open("rnaseq_rawcounts.txt") as handle:
            text = io.TextIOWrapper(handle)
            rna_header = text.readline().rstrip("\n").split("\t")
            cldn4 = None
            for line in text:
                parts = line.rstrip("\n").split("\t")
                if parts[0] == "CLDN4":
                    cldn4 = [float(x) for x in parts[1:]]
                    break
    if cldn4 is None:
        raise ValueError("Liu CLDN4 missing")
    expression = dict(zip(rna_header[1:], cldn4, strict=True))
    sample_rows = []
    for patient, count in expression.items():
        clin = clinical[patient]
        recist = str(clin["BR"])
        group = (
            "Responder"
            if recist in {"CR", "PR"}
            else "Nonresponder"
            if recist in {"SD", "PD"}
            else ""
        )
        sample_rows.append(
            {
                "cohort": "Liu2019_NatMed",
                "sample": patient,
                "patient": patient,
                "visit": "Pre",
                "response_raw": recist,
                "primary_type": clin["Primary_Type"],
                "treatment": clin["Tx"],
                "prior_ctla4": clin["priorCTLA4"],
                "biopsy_context": clin["biopsyContext (1=Pre-Ipi; 2=On-Ipi; 3=Pre-PD1; 4=On-PD1)"],
                "cldn4_value": count,
                "cldn4_log2": math.log2(count + 1),
                "binary_group": group,
            }
        )
    eligible = [
        row
        for row in sample_rows
        if row["binary_group"] and row["biopsy_context"] == 3
    ]
    values = np.array([row["cldn4_log2"] for row in eligible])
    responder = np.array([row["binary_group"] == "Responder" for row in eligible])
    return sample_rows, summarize_expression(values, responder)


def prat_inventory() -> dict:
    rows = parse_geo_sample_rows(DATA / "GSE93157_series_matrix.txt.gz")
    source = geo_one(rows, "!Sample_source_name_ch1")
    chars = geo_characteristics(rows)
    melanoma = sum(item.upper() == "MELANOMA" for item in source)
    cldn4 = False
    with gzip.open(DATA / "GSE93157_series_matrix.txt.gz", "rt") as handle:
        for line in handle:
            if line.startswith("ID_REF") or line.startswith('"ID_REF"'):
                for rest in handle:
                    if "CLDN4" in rest.upper():
                        cldn4 = True
                        break
                break
    return {
        "cohort": "Prat2017_GSE93157",
        "status": "skipped",
        "reason": (
            f"Melanoma n={melanoma} is on the nCounter PanCancer Immune panel; "
            f"CLDN4 is {'present' if cldn4 else 'absent'} from the public matrix."
        ),
        "n_melanoma": melanoma,
        "response_field": ",".join(sorted(set(chars["best.resp"]))),
    }


def make_riaz_figure(
    baseline: dict[str, np.ndarray],
    on_treatment: dict[str, np.ndarray],
    changes: dict[str, np.ndarray],
    summaries: dict[str, dict],
) -> None:
    plt.rcParams["svg.hashsalt"] = "GSE91061-CLDN4"
    plt.style.use("seaborn-v0_8-whitegrid")
    fig, axes = plt.subplots(1, 3, figsize=(12, 4.1))
    rng = np.random.default_rng(SEED)
    panels = [
        ("Pretreatment (primary)", baseline, "log2(FPKM + 1)", "pretreatment"),
        ("On-treatment (exploratory)", on_treatment, "log2(FPKM + 1)", "on_treatment"),
        ("Paired on − pre (exploratory)", changes, "Δ log2(FPKM + 1)", "paired_change"),
    ]
    colors = ["#6b7280", "#2563eb"]
    for ax, (title, groups, ylabel, key) in zip(axes, panels, strict=True):
        ordered = [groups["Nonresponder"], groups["Responder"]]
        bp = ax.boxplot(
            ordered,
            positions=[1, 2],
            widths=0.55,
            patch_artist=True,
            showfliers=False,
            medianprops={"color": "black", "linewidth": 1.5},
        )
        for box, color in zip(bp["boxes"], colors, strict=True):
            box.set_facecolor(color)
            box.set_alpha(0.35)
        for pos, values, color in zip([1, 2], ordered, colors, strict=True):
            jitter = rng.uniform(-0.12, 0.12, len(values))
            ax.scatter(
                pos + jitter,
                values,
                s=23,
                color=color,
                alpha=0.8,
                edgecolor="white",
                linewidth=0.35,
                zorder=3,
            )
        ax.set_xticks([1, 2], [f"NR\nn={len(ordered[0])}", f"R\nn={len(ordered[1])}"])
        ax.set_title(title, fontsize=10, weight="bold")
        ax.set_ylabel(ylabel)
        ax.text(
            0.5,
            0.98,
            f"MW p={summaries[key]['mann_whitney_p']:.3g}",
            transform=ax.transAxes,
            ha="center",
            va="top",
            fontsize=9,
        )
    fig.suptitle("Riaz 2017 melanoma: CLDN4 expression versus nivolumab response", y=1.02)
    fig.tight_layout()
    fig.savefig(ROOT / "cldn4_vs_response.png", dpi=220, bbox_inches="tight")
    fig.savefig(ROOT / "cldn4_vs_response.svg", bbox_inches="tight", metadata={"Date": None})
    plt.close(fig)


def make_forest(cohorts: list[dict], pooled: dict[str, float]) -> None:
    plt.rcParams["svg.hashsalt"] = "MELANOMA-CLDN4-OR"
    plt.style.use("seaborn-v0_8-whitegrid")
    labels = [row["label"] for row in cohorts] + ["Mantel–Haenszel (Riaz+Hugo+Liu)"]
    ors = [row["or_cldn4high_response"] for row in cohorts] + [pooled["or_mh"]]
    lo = [row["or_ci_low"] for row in cohorts] + [pooled["or_ci_low"]]
    hi = [row["or_ci_high"] for row in cohorts] + [pooled["or_ci_high"]]
    fig, ax = plt.subplots(figsize=(8.6, 3.8))
    y = np.arange(len(labels))[::-1]
    ax.axvline(1.0, color="black", linewidth=1)
    ax.errorbar(
        ors,
        y,
        xerr=[np.array(ors) - np.array(lo), np.array(hi) - np.array(ors)],
        fmt="o",
        color="#111827",
        ecolor="#4b5563",
        capsize=3,
    )
    ax.set_yticks(y, labels)
    ax.set_xscale("log")
    ax.set_xlabel("Odds ratio for response in CLDN4-high vs CLDN4-low (Haldane–Anscombe)")
    ax.set_title("Open melanoma ICI RNA: CLDN4-high versus response")
    ax.set_xlim(0.05, 20)
    fig.tight_layout()
    fig.savefig(ROOT / "open_melanoma_forest.png", dpi=220, bbox_inches="tight")
    fig.savefig(ROOT / "open_melanoma_forest.svg", bbox_inches="tight", metadata={"Date": None})
    plt.close(fig)


def main() -> None:
    download_inputs()
    riaz_samples, riaz_primary, riaz_groups = load_riaz()
    riaz_changes, riaz_extra, riaz_extra_groups = riaz_exploratory(riaz_samples)
    hugo_samples, hugo_summary = load_hugo()
    aus_samples, aus_collapsed, aus_summary = load_auslander()
    liu_samples, liu_summary = load_liu()
    prat = prat_inventory()

    riaz_summaries = {**riaz_primary, **riaz_extra}
    write_csv(
        ROOT / "summary.csv",
        [{"analysis": name, **values} for name, values in riaz_summaries.items()],
    )
    write_csv(ROOT / "sample_level.csv", riaz_samples)
    write_csv(ROOT / "paired_changes.csv", riaz_changes)
    write_csv(ROOT / "hugo_sample_level.csv", hugo_samples)
    write_csv(ROOT / "auslander_sample_level.csv", aus_samples)
    write_csv(ROOT / "auslander_patient_level.csv", aus_collapsed)
    write_csv(ROOT / "liu_sample_level.csv", liu_samples)

    cohort_rows = [
        {
            "cohort": "Riaz2017_GSE91061",
            "label": "Riaz 2017 nivo (pre)",
            "role": "primary_analog",
            "n_total": riaz_primary["pretreatment"]["n_responder"]
            + riaz_primary["pretreatment"]["n_nonresponder"],
            **riaz_primary["pretreatment"],
        },
        {
            "cohort": "Hugo2016_GSE78220",
            "label": "Hugo 2016 anti–PD-1 (pre)",
            "role": "open_melanoma",
            "n_total": hugo_summary["n_responder"] + hugo_summary["n_nonresponder"],
            **hugo_summary,
        },
        {
            "cohort": "Liu2019_NatMed",
            "label": "Liu 2019 anti–PD-1 (pre-PD1 RNA)",
            "role": "open_melanoma",
            "n_total": liu_summary["n_responder"] + liu_summary["n_nonresponder"],
            **liu_summary,
        },
        {
            "cohort": "Auslander2018_GSE115821",
            "label": "Auslander 2018 ICI (pre, patient)",
            "role": "open_melanoma_underpowered",
            "n_total": aus_summary["n_responder"] + aus_summary["n_nonresponder"],
            **aus_summary,
        },
    ]
    write_csv(ROOT / "open_melanoma_summary.csv", cohort_rows)
    pooled_all = mantel_haenszel(cohort_rows)
    pooled_main = mantel_haenszel(cohort_rows[:3])
    write_csv(
        ROOT / "open_melanoma_pooled.csv",
        [
            {"pool": "Riaz+Hugo+Liu", **pooled_main},
            {"pool": "Riaz+Hugo+Liu+Auslander", **pooled_all},
        ],
    )
    write_csv(
        ROOT / "inventory.csv",
        [
            {
                "cohort": "Riaz2017_GSE91061",
                "status": "analyzed",
                "reason": "Public FPKM + GEO PRCR/SD/PD. Primary pretreatment analog.",
            },
            {
                "cohort": "Hugo2016_GSE78220",
                "status": "analyzed",
                "reason": "Public FPKM + GEO anti-PD-1 RECIST.",
            },
            {
                "cohort": "Liu2019_NatMed",
                "status": "analyzed",
                "reason": "Public raw counts (GitHub addData) + Nature supplementary Table 1 response.",
            },
            {
                "cohort": "Auslander2018_GSE115821",
                "status": "analyzed_underpowered",
                "reason": "Public counts + R/NR, but only two pretreatment responders after patient collapse.",
            },
            prat,
            {
                "cohort": "Gide2019_PRJEB23709",
                "status": "skipped",
                "reason": "ENA raw RNA only; no public processed matrix with response labels.",
            },
            {
                "cohort": "VanAllen2015",
                "status": "skipped",
                "reason": "RNA is controlled-access; not reconstructed here.",
            },
            {
                "cohort": "Chen2016_GSE67501",
                "status": "skipped",
                "reason": "Public anti-PD-1 RNA is RCC, not melanoma.",
            },
        ],
    )

    make_riaz_figure(
        riaz_groups,
        riaz_extra_groups["on_treatment"],
        riaz_extra_groups["paired_change"],
        riaz_summaries,
    )
    make_forest(cohort_rows[:3], pooled_main)

    with (ROOT / "provenance.tsv").open("w") as handle:
        handle.write("file\tsha256\turl\n")
        for name, url in URLS.items():
            handle.write(f"{name}\t{sha256(DATA / name)}\t{url}\n")

    print("Riaz GEO responses:", dict(Counter(row["response_raw"] for row in riaz_samples)))
    for row in cohort_rows:
        print(
            f"{row['cohort']}: R={row['n_responder']} NR={row['n_nonresponder']} "
            f"AUC={row['auc_responder_higher']:.3f} MW p={row['mann_whitney_p']:.4g} "
            f"OR_high={row['or_cldn4high_response']:.3f} Fisher p={row['or_fisher_p']:.4g}"
        )
    print(
        f"MH Riaz+Hugo+Liu OR={pooled_main['or_mh']:.3f} "
        f"({pooled_main['or_ci_low']:.3f}-{pooled_main['or_ci_high']:.3f})"
    )


if __name__ == "__main__":
    main()
