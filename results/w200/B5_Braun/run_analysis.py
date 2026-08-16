#!/usr/bin/env python3
"""Reproduce the Braun 2020 CLDN4-versus-objective-response analysis."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import urllib.request
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.stats import fisher_exact, mannwhitneyu, norm


SOURCE_URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1038%2Fs41591-020-0839-y/MediaObjects/"
    "41591_2020_839_MOESM2_ESM.xlsx"
)
SOURCE_SHA256 = "a1f8683676d416b255291dcaa007aa0fa48c736c23e6a7a566664ba528bf205b"
RESPONDER_LABELS = {"CR", "PR", "CRPR"}
NONRESPONDER_LABELS = {"SD", "PD"}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def obtain_workbook(path: Path) -> None:
    if path.exists() and sha256(path) == SOURCE_SHA256:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".download")
    print(f"Downloading public supplementary workbook to {path} ...")
    urllib.request.urlretrieve(SOURCE_URL, temporary)
    observed = sha256(temporary)
    if observed != SOURCE_SHA256:
        temporary.unlink(missing_ok=True)
        raise RuntimeError(
            f"Source checksum changed: expected {SOURCE_SHA256}, observed {observed}"
        )
    temporary.replace(path)


def read_records(workbook_path: Path) -> tuple[list[dict], dict]:
    workbook = openpyxl.load_workbook(
        workbook_path, read_only=True, data_only=True
    )

    clinical_sheet = workbook["S1_Clinical_and_Immune_Data"]
    clinical_rows = clinical_sheet.iter_rows(min_row=2, values_only=True)
    clinical_header = next(clinical_rows)
    clinical_index = {name: index for index, name in enumerate(clinical_header)}
    required_fields = {"SUBJID", "Cohort", "Arm", "RNA_ID", "ORR"}
    if not required_fields.issubset(clinical_index):
        raise RuntimeError("Expected clinical fields are absent from source workbook")

    clinical_by_rna = {}
    for row in clinical_rows:
        rna_id = row[clinical_index["RNA_ID"]]
        if rna_id not in (None, "NA"):
            if rna_id in clinical_by_rna:
                raise RuntimeError(f"Duplicate clinical RNA_ID: {rna_id}")
            clinical_by_rna[rna_id] = {
                field: row[clinical_index[field]] for field in required_fields
            }

    expression_sheet = workbook["S4A_RNA_Expression"]
    expression_rows = expression_sheet.iter_rows(min_row=2, values_only=True)
    expression_header = next(expression_rows)
    cldn4_rows = [row for row in expression_rows if row[0] == "CLDN4"]
    if len(cldn4_rows) != 1:
        raise RuntimeError(f"Expected one CLDN4 row; found {len(cldn4_rows)}")
    cldn4_by_rna = dict(zip(expression_header[1:], cldn4_rows[0][1:]))

    if set(cldn4_by_rna) != set(clinical_by_rna):
        raise RuntimeError("Clinical and expression RNA identifiers do not match")

    records = []
    exclusion_counts = {
        "non_nivolumab_arm": 0,
        "objective_response_not_evaluable": 0,
    }
    for rna_id, expression in cldn4_by_rna.items():
        clinical = clinical_by_rna[rna_id]
        if clinical["Arm"] != "NIVOLUMAB":
            exclusion_counts["non_nivolumab_arm"] += 1
            continue
        response_label = clinical["ORR"]
        if response_label in RESPONDER_LABELS:
            response = "Responder (CR/PR)"
            response_binary = 1
        elif response_label in NONRESPONDER_LABELS:
            response = "Nonresponder (SD/PD)"
            response_binary = 0
        else:
            exclusion_counts["objective_response_not_evaluable"] += 1
            continue
        records.append(
            {
                "subject_id": clinical["SUBJID"],
                "rna_id": rna_id,
                "cohort": clinical["Cohort"],
                "arm": clinical["Arm"],
                "orr_source_label": response_label,
                "response": response,
                "response_binary": response_binary,
                "cldn4_normalized_expression": float(expression),
            }
        )

    workbook.close()
    if len({record["subject_id"] for record in records}) != len(records):
        raise RuntimeError("Analysis set unexpectedly contains repeated subjects")
    records.sort(key=lambda record: (record["cohort"], record["subject_id"]))
    return records, exclusion_counts


def quantiles(values: np.ndarray) -> dict:
    q1, median, q3 = np.quantile(values, [0.25, 0.5, 0.75])
    return {
        "n": int(values.size),
        "median": float(median),
        "q1": float(q1),
        "q3": float(q3),
    }


def bootstrap_median_difference(
    responder: np.ndarray, nonresponder: np.ndarray
) -> tuple[float, float]:
    rng = np.random.default_rng(20260816)
    differences = np.empty(20_000)
    for index in range(differences.size):
        differences[index] = np.median(
            rng.choice(responder, responder.size, replace=True)
        ) - np.median(rng.choice(nonresponder, nonresponder.size, replace=True))
    lower, upper = np.quantile(differences, [0.025, 0.975])
    return float(lower), float(upper)


def woolf_odds_ratio(high_resp: int, high_non: int, low_resp: int, low_non: int) -> dict:
    if min(high_resp, high_non, low_resp, low_non) == 0:
        raise RuntimeError("Odds ratio is undefined for a zero cell")
    odds_ratio = (high_resp / high_non) / (low_resp / low_non)
    log_or = np.log(odds_ratio)
    se = np.sqrt(1 / high_resp + 1 / high_non + 1 / low_resp + 1 / low_non)
    return {
        "odds_ratio": float(odds_ratio),
        "odds_ratio_95_ci": [
            float(np.exp(log_or - 1.96 * se)),
            float(np.exp(log_or + 1.96 * se)),
        ],
    }


def fit_logistic(x: np.ndarray, y: np.ndarray) -> dict:
    design = np.column_stack([np.ones(x.size), x])
    beta = np.zeros(2)
    for _ in range(50):
        linear = design @ beta
        linear = np.clip(linear, -30, 30)
        fitted = 1.0 / (1.0 + np.exp(-linear))
        weights = fitted * (1.0 - fitted)
        information = design.T @ (weights[:, None] * design)
        score = design.T @ (y - fitted)
        try:
            step = np.linalg.solve(information, score)
        except np.linalg.LinAlgError as exc:
            raise RuntimeError("Logistic information matrix is singular") from exc
        beta = beta + step
        if np.max(np.abs(step)) < 1e-10:
            break
    se = np.sqrt(np.diag(np.linalg.inv(information)))
    odds_ratio = float(np.exp(beta[1]))
    return {
        "log_odds_per_sd": float(beta[1]),
        "odds_ratio_per_sd": odds_ratio,
        "odds_ratio_per_sd_95_ci": [
            float(np.exp(beta[1] - 1.96 * se[1])),
            float(np.exp(beta[1] + 1.96 * se[1])),
        ],
        "wald_p_value_two_sided": float(2 * (1 - norm.cdf(abs(beta[1] / se[1])))),
    }


def calculate_results(records: list[dict], exclusion_counts: dict) -> dict:
    responder = np.array(
        [
            record["cldn4_normalized_expression"]
            for record in records
            if record["response_binary"] == 1
        ]
    )
    nonresponder = np.array(
        [
            record["cldn4_normalized_expression"]
            for record in records
            if record["response_binary"] == 0
        ]
    )
    test = mannwhitneyu(
        responder, nonresponder, alternative="two-sided", method="asymptotic"
    )
    median_difference = float(np.median(responder) - np.median(nonresponder))
    ci_lower, ci_upper = bootstrap_median_difference(responder, nonresponder)

    expression = np.array(
        [record["cldn4_normalized_expression"] for record in records]
    )
    response = np.array([record["response_binary"] for record in records])
    split = float(np.median(expression))
    high = expression >= split
    high_resp = int((high & (response == 1)).sum())
    high_non = int((high & (response == 0)).sum())
    low_resp = int((~high & (response == 1)).sum())
    low_non = int((~high & (response == 0)).sum())
    fisher = fisher_exact(
        [[low_resp, low_non], [high_resp, high_non]], alternative="two-sided"
    )
    dichotomized = {
        "cut": "pre-specified median of included CLDN4 values; high = at or above median",
        "median_cut": split,
        "n": len(records),
        "n_high": int(high.sum()),
        "n_low": int((~high).sum()),
        "high_responder": high_resp,
        "high_nonresponder": high_non,
        "low_responder": low_resp,
        "low_nonresponder": low_non,
        **woolf_odds_ratio(high_resp, high_non, low_resp, low_non),
        "fisher_p_value_two_sided": float(fisher.pvalue),
        "odds_ratio_definition": (
            "odds of CR/PR in CLDN4-high versus CLDN4-low"
        ),
    }
    z_expression = (expression - expression.mean()) / expression.std(ddof=1)
    continuous_or = fit_logistic(z_expression, response.astype(float))
    headline = {
        "cohort": "open RCC ICI (Braun 2020 CheckMate 009/010/025 nivolumab RNA)",
        "endpoint": "objective response CR/PR vs SD/PD",
        "exposure": "CLDN4 high vs low (median split)",
        "OR": dichotomized["odds_ratio"],
        "OR_95_ci": dichotomized["odds_ratio_95_ci"],
        "n": dichotomized["n"],
        "p": dichotomized["fisher_p_value_two_sided"],
        "p_test": "two-sided Fisher exact",
    }

    cohort_results = []
    for cohort in sorted({record["cohort"] for record in records}):
        cohort_responder = np.array(
            [
                record["cldn4_normalized_expression"]
                for record in records
                if record["cohort"] == cohort and record["response_binary"] == 1
            ]
        )
        cohort_nonresponder = np.array(
            [
                record["cldn4_normalized_expression"]
                for record in records
                if record["cohort"] == cohort and record["response_binary"] == 0
            ]
        )
        cohort_test = mannwhitneyu(
            cohort_responder,
            cohort_nonresponder,
            alternative="two-sided",
            method="asymptotic",
        )
        cohort_results.append(
            {
                "cohort": cohort,
                "responder": quantiles(cohort_responder),
                "nonresponder": quantiles(cohort_nonresponder),
                "mann_whitney_u": float(cohort_test.statistic),
                "p_value_two_sided": float(cohort_test.pvalue),
                "common_language_effect_probability": float(
                    cohort_test.statistic
                    / (cohort_responder.size * cohort_nonresponder.size)
                ),
            }
        )

    return {
        "source": {
            "citation": (
                "Braun DA et al. Nature Medicine 2020;26:909-918. "
                "doi:10.1038/s41591-020-0839-y"
            ),
            "supplement_url": SOURCE_URL,
            "supplement_sha256": SOURCE_SHA256,
            "clinical_sheet": "S1_Clinical_and_Immune_Data",
            "expression_sheet": "S4A_RNA_Expression",
        },
        "analysis_set": {
            "included_n": len(records),
            "definition": (
                "Nivolumab-treated, RNA-profiled subjects with evaluable ORR; "
                "responder=CR/PR (including pooled source label CRPR), "
                "nonresponder=SD/PD"
            ),
            "excluded": exclusion_counts,
        },
        "headline_or_n_p": headline,
        "pooled_result": {
            "responder": quantiles(responder),
            "nonresponder": quantiles(nonresponder),
            "median_difference_responder_minus_nonresponder": median_difference,
            "median_difference_bootstrap_95_ci": [ci_lower, ci_upper],
            "mann_whitney_u": float(test.statistic),
            "p_value_two_sided": float(test.pvalue),
            "common_language_effect_probability": float(
                test.statistic / (responder.size * nonresponder.size)
            ),
            "median_split_odds_ratio": dichotomized,
            "logistic_odds_ratio_per_sd": continuous_or,
        },
        "cohort_descriptive_results": cohort_results,
    }


def write_samples(records: list[dict], output_path: Path) -> None:
    with output_path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)


def stable_jitter(record: dict) -> float:
    digest = hashlib.sha256(record["rna_id"].encode()).digest()
    return (int.from_bytes(digest[:4], "big") / (2**32 - 1) - 0.5) * 0.30


def make_figure(records: list[dict], results: dict, output_path: Path) -> None:
    groups = [
        ("Nonresponder (SD/PD)", 0, "#68788A"),
        ("Responder (CR/PR)", 1, "#C65A46"),
    ]
    fig, axis = plt.subplots(figsize=(5.2, 4.8))
    values = []
    for label, binary, color in groups:
        selected = [
            record for record in records if record["response_binary"] == binary
        ]
        group_values = [
            record["cldn4_normalized_expression"] for record in selected
        ]
        values.append(group_values)
        x = [binary + stable_jitter(record) for record in selected]
        axis.scatter(
            x,
            group_values,
            s=20,
            alpha=0.62,
            color=color,
            edgecolors="white",
            linewidths=0.35,
            zorder=3,
        )
    box = axis.boxplot(
        values,
        positions=[0, 1],
        widths=0.48,
        patch_artist=True,
        showfliers=False,
        medianprops={"color": "#111111", "linewidth": 1.6},
        whiskerprops={"color": "#333333"},
        capprops={"color": "#333333"},
    )
    for patch, (_, _, color) in zip(box["boxes"], groups):
        patch.set_facecolor(color)
        patch.set_alpha(0.18)
        patch.set_edgecolor(color)
    headline = results["headline_or_n_p"]
    axis.set_title(
        (
            "Open RCC ICI (Braun 2020): CLDN4 vs objective response\n"
            f"OR={headline['OR']:.2f} "
            f"({headline['OR_95_ci'][0]:.2f}–{headline['OR_95_ci'][1]:.2f}); "
            f"n={headline['n']}; "
            f"p={headline['p']:.2f}"
        ),
        loc="left",
        fontsize=12,
        linespacing=1.45,
    )
    axis.set_ylabel("CLDN4 normalized RNA expression")
    axis.set_xticks(
        [0, 1],
        [
            f"SD/PD\nn={len(values[0])}",
            f"CR/PR\nn={len(values[1])}",
        ],
    )
    axis.spines[["top", "right"]].set_visible(False)
    axis.grid(axis="y", color="#E5E5E5", linewidth=0.8)
    axis.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(output_path, dpi=180)
    plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path(".cache/Braun2020_supplementary_tables.xlsx"),
        help="Source workbook path; downloaded and checksum-verified if absent",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parent,
    )
    args = parser.parse_args()

    obtain_workbook(args.workbook)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    records, exclusion_counts = read_records(args.workbook)
    results = calculate_results(records, exclusion_counts)
    write_samples(records, args.output_dir / "analysis_samples.csv")
    with (args.output_dir / "results.json").open("w") as handle:
        json.dump(results, handle, indent=2)
        handle.write("\n")
    headline = results["headline_or_n_p"]
    with (args.output_dir / "or_n_p.json").open("w") as handle:
        json.dump(headline, handle, indent=2)
        handle.write("\n")
    with (args.output_dir / "or_n_p.tsv").open("w", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["cohort", "endpoint", "exposure", "OR", "OR_lo", "OR_hi", "n", "p", "p_test"],
            delimiter="\t",
        )
        writer.writeheader()
        writer.writerow(
            {
                "cohort": headline["cohort"],
                "endpoint": headline["endpoint"],
                "exposure": headline["exposure"],
                "OR": f"{headline['OR']:.6f}",
                "OR_lo": f"{headline['OR_95_ci'][0]:.6f}",
                "OR_hi": f"{headline['OR_95_ci'][1]:.6f}",
                "n": headline["n"],
                "p": f"{headline['p']:.6f}",
                "p_test": headline["p_test"],
            }
        )
    make_figure(records, results, args.output_dir / "CLDN4_vs_ORR.png")
    print(json.dumps(headline, indent=2))


if __name__ == "__main__":
    main()
